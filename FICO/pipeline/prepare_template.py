"""
Stamp INPUT-ONLY Named Ranges onto a combined CFI template.

Copies pristine CFI 3-statement + DCF originals into
FICO/templates/CFI_Template.xlsx and registers Defined Names solely for
hardcoded input cells (never formula-output cells).
"""

from __future__ import annotations

import shutil
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

from .named_range_map import SHEET_3S, SHEET_DCF, TEMPLATE_NAMED_RANGES

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parent
ORIG_3S = REPO / "DCF resources" / "originals" / "CFI_3-Statement-Model-Complete.xlsx"
ORIG_DCF = REPO / "DCF resources" / "originals" / "CFI_DCF-Model.xlsx"
OUT_TEMPLATE = ROOT / "templates" / "CFI_Template.xlsx"


def _copy_sheet(src_ws, dst_wb, title: str):
    """Copy a worksheet (values/styles/merged cells) into dst_wb."""
    dst = dst_wb.create_sheet(title)
    for row in src_ws.iter_rows():
        for cell in row:
            nc = dst.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                nc.font = copy(cell.font)
                nc.border = copy(cell.border)
                nc.fill = copy(cell.fill)
                nc.number_format = cell.number_format
                nc.protection = copy(cell.protection)
                nc.alignment = copy(cell.alignment)
    for merged in src_ws.merged_cells.ranges:
        dst.merge_cells(str(merged))
    for col_letter, col_dim in src_ws.column_dimensions.items():
        dst.column_dimensions[col_letter].width = col_dim.width
    for row_idx, row_dim in src_ws.row_dimensions.items():
        dst.row_dimensions[row_idx].height = row_dim.height
    return dst


def build_template(
    out_path: Path = OUT_TEMPLATE,
    *,
    orig_3s: Path = ORIG_3S,
    orig_dcf: Path = ORIG_DCF,
) -> Path:
    if not orig_3s.exists():
        raise FileNotFoundError(f"Missing pristine 3-statement template: {orig_3s}")
    if not orig_dcf.exists():
        raise FileNotFoundError(f"Missing pristine DCF template: {orig_dcf}")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    # Start from 3-statement workbook (keeps Cover Page + 3 Statement Model)
    # IMPORTANT: data_only=False (default) — preserve formulas, not cached values.
    wb = load_workbook(orig_3s, data_only=False)
    dcf_wb = load_workbook(orig_dcf, data_only=False)

    # Remove CapIQ junk names we don't need; keep workbook clean for our ranges
    for name in list(wb.defined_names.keys()):
        if name.startswith("IQ_") or name in {
            "CIQWBGuid",
            "asd",
            "Forecast",
            "Step_1",
            "Step_2",
            "Step_3",
            "Step_4",
            "Step_5",
            "Step_6",
            "Step1",
            "Step2",
            "Step3",
            "Step4",
            "Step5",
            "Step6",
        }:
            del wb.defined_names[name]

    # Add DCF Model sheet if missing
    if SHEET_DCF not in wb.sheetnames:
        _copy_sheet(dcf_wb[SHEET_DCF], wb, SHEET_DCF)

    # Ensure expected sheets exist
    if SHEET_3S not in wb.sheetnames:
        raise RuntimeError(f"Template missing sheet {SHEET_3S!r}")

    # Annotate cover
    if "Cover Page" in wb.sheetnames:
        cover = wb["Cover Page"]
        cover["C12"] = "FICO — CFI Template (Named-Range Inputs)"
        cover["C20"] = (
            "Named Ranges mark INPUT cells only. Formula cells (Gross Profit, NI, UFCF, EV) "
            "have NO names and must never be overwritten. Open in Excel to recalculate."
        )

    # Register / replace our input Named Ranges
    for name, attr in TEMPLATE_NAMED_RANGES.items():
        if name in wb.defined_names:
            del wb.defined_names[name]
        wb.defined_names.add(DefinedName(name=name, attr_text=attr))

    # Widen year columns so large $000s figures do not render as ########
    if SHEET_3S in wb.sheetnames:
        ws = wb[SHEET_3S]
        ws.column_dimensions["B"].width = 42
        for col in range(4, 15):
            ws.column_dimensions[get_column_letter(col)].width = 16
    if SHEET_DCF in wb.sheetnames:
        ws = wb[SHEET_DCF]
        ws.column_dimensions["B"].width = 28
        for col in range(3, 12):
            ws.column_dimensions[get_column_letter(col)].width = 16

    wb.save(out_path)
    print(f"Wrote template with {len(TEMPLATE_NAMED_RANGES)} input Named Ranges → {out_path}")
    return out_path


def main() -> int:
    build_template()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
