"""Plain-English explanations for every 3-statement assumption row (MODEL6).

SOURCE fields always include a clickable URL (col U) so users can open the filing/data.
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import List, Tuple

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

from .named_range_map import SHEET_3S

_FORECAST_COLS = ("J", "K", "L", "M", "N")
_COMMENT_AUTHOR = "vengeanceaiUSCMODEL6"

# Canonical source URLs (clickable in Excel col U)
URL_10K = "https://www.sec.gov/Archives/edgar/data/814547/000081454725000030/fico-20250930.htm"
URL_10Q = "https://www.sec.gov/Archives/edgar/data/814547/000081454726000030/fico-20260630.htm"
URL_GUIDANCE = (
    "https://www.sec.gov/Archives/edgar/data/814547/000081454726000031/exhibit991erq32026.htm"
)
URL_FACTS = "https://data.sec.gov/api/xbrl/companyfacts/CIK0000814547.json"
URL_8K_NOTES = "https://www.sec.gov/Archives/edgar/data/814547/000119312526117936/d56220d8k.htm"
URL_DAMODARAN_FCFF = "https://pages.stern.nyu.edu/~adamodar/New_Home_Page/datafile/histfcff.html"
URL_IR = "https://www.fico.com/en/investors"

LINK_FONT = Font(name="Calibri", size=9, color="0563C1", underline="single")

# (row, short_name, what_it_is, how_set, why, source_label, source_url)
ASSUMPTION_EXPLANATIONS: List[Tuple[int, str, str, str, str, str, str]] = [
    (
        7,
        "Revenue Growth",
        "Year-over-year % increase in revenue for each forecast year.",
        "Yellow POLICY input: 27% / 16% / 13% / 10% / 7%.",
        "Y1 ≈ company FY2026 revenue guidance (~$2.53B / FY25 $1.991B − 1 ≈ 27%). "
        "Then fade toward terminal g=3% (not straight-lined). This is the main "
        "forward-looking judgment; it drives Rev → GP → EBT → Net Earnings → CF.",
        "SEC EX-99.1 Q3 FY2026 — updated FY2026 revenue guidance $2.53B",
        URL_GUIDANCE,
    ),
    (
        8,
        "COGS % of Revenue",
        "Cost of revenues as a % of sales (gross margin = 1 − this).",
        "Excel equation: I25/I24 (FY25 COGS ÷ FY25 Revenue) held flat.",
        "Locks in latest reported cost structure (~17.8%). Scores mix is high-margin; "
        "holding FY25 is conservative vs further mix shift.",
        "SEC 10-K FY2025 — Cost of revenues / Total revenues",
        URL_10K,
    ),
    (
        9,
        "SGA % of Revenue",
        "Operating opex (Salaries & Benefits / SG&A) as % of sales.",
        "Equation: MAX(20%, (I28 − 10,922)/I24 − 25bps × year). "
        "Strips FY25 restructuring; then −0.25% of sales each year.",
        "MODEL6 fix: enterprise software needs lasting G&A (sales, legal, compliance). "
        "Floor raised from 5%→20%; grind cut from 50→25 bps so SGA stays in a "
        "realistic ~22–25% band near FY25 normalized (~25%).",
        "SEC 10-K FY2025 — SG&A + restructuring note",
        URL_10K,
    ),
    (
        10,
        "R&D % of Revenue",
        "Research & development (template label may say Rent and Overhead) as % of sales.",
        "Excel equation: I29/I24 held flat (~9.46%).",
        "FICO reinvests steadily in analytics/software. Flat % grows dollars with revenue.",
        "SEC 10-K FY2025 — Research and development",
        URL_10K,
    ),
    (
        11,
        "D&A % of Avg PP&E",
        "Depreciation rate from FY25; dollars applied to average PP&E so new CapEx is depreciated.",
        "Rate = I30/I44 (FY25 DA ÷ FY25 PPE). "
        "Forecast DA$ = ((Open + Open+CapEx)/2) × rate.",
        "MODEL6 fix: prior model depreciated only opening PP&E, so ~$96M of forecast "
        "CapEx never hit D&A/FCFF add-back. Average PP&E (Open ↔ pre-DA close) "
        "depreciates additions in the year they are placed in service.",
        "SEC 10-K FY2025 — D&A and PP&E",
        URL_10K,
    ),
    (
        12,
        "Interest % of Debt Open",
        "Interest expense as % of opening debt balance.",
        "Excel equation: I101/I98 (FY25 interest ÷ FY25 opening debt in schedule).",
        "Approximates average coupon / cost of debt on the book. Debt held flat "
        "(issuance = 0), so interest stays linked to that stock.",
        "SEC 10-K FY2025 — Interest expense; see also 8-K 6.250% notes",
        URL_10K,
    ),
    (
        13,
        "Tax Rate (% of EBT)",
        "Effective book tax rate applied to forecast EBT.",
        "Excel equation: I35/I33 (FY25 tax ÷ simplified FY25 EBT in this template).",
        "Uses reported effective rate (~19% on template EBT; ~18.8% on 10-K EBT). "
        "UFCF in DCF also uses unlevered EBIT × t.",
        "SEC 10-K FY2025 — Provision for income taxes / Income before taxes",
        URL_10K,
    ),
    (
        15,
        "Accounts Receivable (Days)",
        "Standard DSO: gross AR days used to project Gross AR = Rev × days/365.",
        "Excel equation: ROUND(I42/I24×365, 0). I42 is GROSS AR (not net of deferred).",
        "MODEL6 fix: DSO must use gross receivables. Deferred Revenue is a contract "
        "liability projected separately (WC row 88 = Rev × FY25 Def/Rev). "
        "NWC = GrossAR + Inv − AP − Deferred.",
        "SEC companyfacts XBRL — AR; DeferredRevenueCurrent separate",
        URL_FACTS,
    ),
    (
        16,
        "Inventory (Days)",
        "Inventory days (Inv = COGS × days/365).",
        "Hard zero — FICO is software / scores; inventory is immaterial.",
        "No inventory cycle to fund.",
        "SEC 10-K FY2025 — Inventory ≈ $0",
        URL_10K,
    ),
    (
        17,
        "Accounts Payable (Days)",
        "AP days used to project AP = COGS × days/365.",
        "Excel equation: ROUND(I48/I25×365, 0) from FY25.",
        "Payable timing offsets AR in operating NWC.",
        "SEC 10-K FY2025 — Accounts payable",
        URL_10K,
    ),
    (
        18,
        "CapEx % of Revenue",
        "Capital investment (PPE + capitalized software) as % of sales.",
        "Equation: fade from I68/I24 (FY25 CapEx/Sales ~1.98%) toward 1.0% steady. "
        "Weights 85%→65%→45%→30%→0% on the peak.",
        "FY25 CapEx was elevated by capitalized internal-use software. Fading avoids "
        "locking a peak reinvestment rate forever.",
        "SEC 10-K FY2025 — PP&E purchases + capitalized software",
        URL_10K,
    ),
    (
        19,
        "Debt Issuance (Repayment)",
        "New borrowing (+) or repayment (−) in the forecast.",
        "POLICY input = 0 every year.",
        "Hold debt stock flat at FY25 level for the explicit period; interest follows. "
        "Net debt for DCF equity bridge uses the later 10-Q amount separately.",
        "SEC 10-Q Q3 FY2026 — debt stock (financing excluded from FCFF)",
        URL_10Q,
    ),
    (
        20,
        "Equity Issued (Repaid)",
        "New equity issued (+) or buybacks (−) in the forecast BS/CF.",
        "POLICY input = 0 every year.",
        "Buybacks are real historically but are financing — excluded from FCFF. "
        "Share count for $/share is the spot shares outstanding assumption.",
        "Damodaran FCFF framework — financing (buybacks) not in FCFF",
        URL_DAMODARAN_FCFF,
    ),
]


def _source_line(label: str, url: str) -> str:
    return f"{label} | {url}"


def _cell_commentary(name: str, how: str, why: str, label: str, url: str) -> str:
    """Text embedded in Excel cell comments (hover on assumption cells)."""
    return (
        f"{name}\n"
        f"HOW: {how}\n"
        f"WHY: {why}\n"
        f"SOURCE: {label}\n"
        f"LINK: {url}"
    )


def _column_c_note(how: str, why: str, label: str, url: str) -> str:
    """Visible note beside the label — must NOT start with '=' (#NAME?)."""
    why_short = why if len(why) <= 120 else why[:117] + "…"
    return f"WHY: {why_short} | SOURCE: {label} | LINK: {url} | HOW: {how}"


def _set_hyperlink(cell, url: str, display: str) -> None:
    cell.value = display
    cell.hyperlink = url
    cell.font = LINK_FONT
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.border = THIN


HEADER_FILL = PatternFill("solid", fgColor="C65911")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
SUB_FILL = PatternFill("solid", fgColor="FCE4D6")
SUB_FONT = Font(name="Calibri", bold=True, size=9, color="833C0C")
BODY_FONT = Font(name="Calibri", size=9)
TITLE_FONT = Font(name="Calibri", bold=True, size=12, color="C65911")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def explanation_rows() -> List[dict]:
    rows = []
    for row, name, what, how, why, label, url in ASSUMPTION_EXPLANATIONS:
        rows.append(
            {
                "row": row,
                "assumption": name,
                "what_it_is": what,
                "how_set_in_model": how,
                "why_this_choice": why,
                "source": label,
                "source_url": url,
            }
        )
    return rows


def write_assumption_explanations(wb) -> None:
    """Put WHY + SOURCE+URL in each assumption cell, col C, and clickable col U."""
    if SHEET_3S not in wb.sheetnames:
        raise RuntimeError(f"Missing sheet {SHEET_3S}")
    ws = wb[SHEET_3S]

    ws["B4"] = (
        "vengeanceaiUSCMODEL6: hover J–N for WHY+SOURCE; click blue Source link in col U "
        "(or LINK: URL in col C / V) to open the filing"
    )
    ws["B4"].font = Font(name="Calibri", bold=True, color="1F4E79")
    ws["B4"].fill = PatternFill("solid", fgColor="D6EAF8")

    for row, name, what, how, why, label, url in ASSUMPTION_EXPLANATIONS:
        ws[f"C{row}"] = _column_c_note(how, why, label, url)
        ws[f"C{row}"].font = Font(name="Calibri", italic=True, size=8, color="595959")
        ws[f"C{row}"].alignment = WRAP

        text = _cell_commentary(name, how, why, label, url)
        for col in _FORECAST_COLS:
            cell = ws[f"{col}{row}"]
            comment = Comment(text, _COMMENT_AUTHOR)
            comment.width = 340
            comment.height = 160
            cell.comment = comment

        label_cell = ws[f"B{row}"]
        label_comment = Comment(text, _COMMENT_AUTHOR)
        label_comment.width = 340
        label_comment.height = 160
        label_cell.comment = label_comment

        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 15, 40)

    # Legend block P–V (U = clickable source, V = raw URL)
    ws["P4"] = "ASSUMPTIONS EXPLAINED — click blue Source (col U) for the filing/data"
    ws["P4"].font = TITLE_FONT
    ws.merge_cells("P4:V4")

    ws["P5"] = (
        "Yellow = policy. Green = FY25-linked equations. "
        "Col U = clickable hyperlink. Col V = full URL (also clickable)."
    )
    ws["P5"].font = Font(name="Calibri", italic=True, size=9, color="595959")
    ws.merge_cells("P5:V5")

    headers = [
        ("P6", "Row"),
        ("Q6", "Assumption"),
        ("R6", "What it is"),
        ("S6", "How set in model"),
        ("T6", "Why this choice"),
        ("U6", "Source (click)"),
        ("V6", "Source URL"),
    ]
    for coord, label in headers:
        cell = ws[coord]
        cell.value = label
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN
        cell.alignment = Alignment(vertical="center")

    for row, name, what, how, why, label, url in ASSUMPTION_EXPLANATIONS:
        for col_idx, val in enumerate([row, name, what, how, why], start=16):  # P–T
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.font = BODY_FONT if col_idx > 17 else SUB_FONT
            cell.alignment = WRAP
            cell.border = THIN
            if col_idx <= 17:
                cell.fill = SUB_FILL

        # U = clickable display label
        u = ws.cell(row=row, column=21)
        _set_hyperlink(u, url, label)
        u.comment = Comment(_cell_commentary(name, how, why, label, url), _COMMENT_AUTHOR)

        # V = raw URL (also hyperlinked for one-click)
        v = ws.cell(row=row, column=22)
        _set_hyperlink(v, url, url)

        # Also hyperlink the row label in B when useful (secondary click target)
        # Keep B as text label; users click U/V.

    ws.column_dimensions["C"].width = 80
    ws.column_dimensions["P"].width = 6
    ws.column_dimensions["Q"].width = 22
    ws.column_dimensions["R"].width = 36
    ws.column_dimensions["S"].width = 42
    ws.column_dimensions["T"].width = 48
    ws.column_dimensions["U"].width = 42
    ws.column_dimensions["V"].width = 55


def export_assumption_explanations_csv(out_dir: Path, *, ticker: str = "FICO") -> Path:
    """Write MODEL6_*_ASSUMPTIONS_EXPLAINED.csv with source_url column."""
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"MODEL6_{ticker}_ASSUMPTIONS_EXPLAINED.csv"
    rows = explanation_rows()
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "row",
                "assumption",
                "what_it_is",
                "how_set_in_model",
                "why_this_choice",
                "source",
                "source_url",
            ],
        )
        writer.writeheader()
        writer.writerows(rows)
    return path
