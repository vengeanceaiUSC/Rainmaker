"""Export separate 3-statement and DCF workbooks (never joined)."""

from __future__ import annotations

from pathlib import Path
from typing import List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from .models import DCFResult, ForecastAssumptions

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


def _style_header(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _write_frame(ws, df: pd.DataFrame, start_row: int = 1, title: str = "") -> int:
    """Write a year-indexed (rows) or metric-indexed frame to a sheet."""
    r = start_row
    if title:
        ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=14)
        r += 2

    # Prefer metrics as rows, years as columns for readability
    if df.index.dtype != object and all(isinstance(i, (int, float)) for i in df.index[:3]):
        out = df.T.copy()
    else:
        out = df.copy()
    out.insert(0, "Line Item", out.index.astype(str))
    out = out.reset_index(drop=True)

    for i, row in enumerate(dataframe_to_rows(out, index=False, header=True)):
        for j, val in enumerate(row, start=1):
            cell = ws.cell(row=r + i, column=j, value=val if not pd.isna(val) else None)
            cell.border = THIN
            if i == 0:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
            elif j > 1 and isinstance(val, (int, float)) and not isinstance(val, bool):
                cell.number_format = "#,##0.0;(#,##0.0);-"
    ws.column_dimensions["A"].width = 36
    for col in range(2, len(out.columns) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    return r + len(out) + 2


def export_three_statement(
    path: Path,
    income: pd.DataFrame,
    balance: pd.DataFrame,
    cashflow: pd.DataFrame,
    assumptions: ForecastAssumptions,
    ticker: str,
    validation_notes: Optional[List[str]] = None,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Income Statement"
    _write_frame(ws, income, title=f"{ticker} — Income Statement ($000s)")

    ws_bs = wb.create_sheet("Balance Sheet")
    _write_frame(ws_bs, balance, title=f"{ticker} — Balance Sheet ($000s)")

    ws_cf = wb.create_sheet("Cash Flow")
    _write_frame(ws_cf, cashflow, title=f"{ticker} — Cash Flow Statement ($000s)")

    ws_a = wb.create_sheet("Assumptions")
    ws_a["A1"] = "Forecast Assumptions (rules-based; no LLM numbers)"
    ws_a["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Revenue growth path", ", ".join(f"{g:.1%}" for g in assumptions.revenue_growth)),
        ("COGS % revenue", assumptions.cogs_pct_revenue),
        ("R&D % revenue", assumptions.rd_pct_revenue),
        ("SG&A % revenue", assumptions.sga_pct_revenue),
        ("D&A % revenue", assumptions.da_pct_revenue),
        ("Tax rate", assumptions.tax_rate),
        ("CapEx % revenue", assumptions.capex_pct_revenue),
        ("NWC % revenue", assumptions.nwc_pct_revenue),
        ("Interest expense level ($000s)", assumptions.interest_expense_level),
        ("Forecast years", assumptions.forecast_years),
        ("WACC (for DCF stage)", assumptions.wacc),
        ("Perpetual growth (for DCF stage)", assumptions.perpetual_growth),
    ]
    ws_a["A3"] = "Parameter"
    ws_a["B3"] = "Value"
    _style_header(ws_a, 3, 2)
    for i, (k, v) in enumerate(rows, start=4):
        ws_a.cell(row=i, column=1, value=k).border = THIN
        cell = ws_a.cell(row=i, column=2, value=v)
        cell.border = THIN
        if isinstance(v, float) and abs(v) <= 2:
            cell.number_format = "0.00%"
        elif isinstance(v, float):
            cell.number_format = "#,##0.0"
    ws_a.column_dimensions["A"].width = 36
    ws_a.column_dimensions["B"].width = 40

    ws_v = wb.create_sheet("Validation")
    ws_v["A1"] = "Historical 3-statement validation"
    ws_v["A1"].font = Font(bold=True, size=14)
    notes = validation_notes or []
    if not notes:
        ws_v["A3"] = "PASS — Assets = Liabilities + Equity; Net Income links to CFS."
    else:
        ws_v["A3"] = "NOTES"
        for i, note in enumerate(notes, start=4):
            ws_v.cell(row=i, column=1, value=note)
    ws_v.column_dimensions["A"].width = 100

    wb.save(path)
    return path


def export_dcf(
    path: Path,
    result: DCFResult,
    annual: pd.DataFrame,
    ticker: str,
    *,
    share_price: float,
    diluted_shares_000s: float,
    cash: float,
    marketable_securities: float,
    total_debt: float,
    net_debt: float,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "DCF Summary"
    ws["A1"] = f"{ticker} — DCF Valuation (Python math; FCFF from 3-statement)"
    ws["A1"].font = Font(bold=True, size=14)

    upside = (result.equity_value_per_share / share_price - 1.0) if share_price else 0.0
    summary = [
        ("Share price (market)", share_price, '"$"#,##0.00'),
        ("Diluted shares (000s)", diluted_shares_000s, "#,##0.000"),
        ("WACC", result.wacc, "0.00%"),
        ("Terminal growth", result.perpetual_growth, "0.00%"),
        ("PV of projected FCFF", sum(result.pv_explicit_fcff), "#,##0.0"),
        ("Terminal value (primary)", result.terminal_value, "#,##0.0"),
        ("PV of terminal value", result.pv_terminal_value, "#,##0.0"),
        ("Enterprise value", result.enterprise_value, "#,##0.0"),
        ("Net debt", net_debt, "#,##0.0"),
        ("Equity value", result.equity_value, "#,##0.0"),
        ("Intrinsic value / share", result.equity_value_per_share, '"$"#,##0.00'),
        ("Upside vs price", upside, "0.00%"),
    ]
    ws["A3"] = "Metric"
    ws["B3"] = "Value"
    _style_header(ws, 3, 2)
    for i, (label, val, fmt) in enumerate(summary, start=4):
        ws.cell(row=i, column=1, value=label).border = THIN
        cell = ws.cell(row=i, column=2, value=val)
        cell.number_format = fmt
        cell.border = THIN
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18

    ws["A18"] = (
        "DCF = Σ CF_t/(1+WACC)^t + TV/(1+WACC)^n  |  "
        "FCFF = EBIT(1-t) + D&A - CapEx - ΔNWC  |  "
        "TV = FCFF_n*(1+g)/(WACC-g)"
    )
    ws["A18"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A18:F18")

    ws_a = wb.create_sheet("Annual FCFF")
    _write_frame(ws_a, annual, title="Projected FCFF bridge and discounting")

    ws_m = wb.create_sheet("Market Inputs")
    ws_m["A1"] = "Market / capital structure inputs (not from LLM)"
    ws_m["A1"].font = Font(bold=True, size=14)
    inputs = [
        ("Cash ($000s)", cash),
        ("Marketable securities ($000s)", marketable_securities),
        ("Total debt ($000s)", total_debt),
        ("Net debt ($000s)", net_debt),
    ]
    ws_m["A3"] = "Input"
    ws_m["B3"] = "Value"
    _style_header(ws_m, 3, 2)
    for i, (k, v) in enumerate(inputs, start=4):
        ws_m.cell(row=i, column=1, value=k).border = THIN
        cell = ws_m.cell(row=i, column=2, value=v)
        cell.border = THIN
        cell.number_format = "#,##0.0"
    ws_m.column_dimensions["A"].width = 36
    ws_m.column_dimensions["B"].width = 18

    ws_n = wb.create_sheet("Notes")
    ws_n["A1"] = "Engine notes"
    ws_n["A1"].font = Font(bold=True, size=14)
    for i, note in enumerate(result.notes, start=3):
        ws_n.cell(row=i, column=1, value=note)
    ws_n.column_dimensions["A"].width = 100

    wb.save(path)
    return path
