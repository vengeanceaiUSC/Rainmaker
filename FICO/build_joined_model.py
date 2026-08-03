#!/usr/bin/env python3
"""
Build a single joined FICO workbook:
  - 03_Three_Statement (CFI template + FICO 10-K history + FICO-scale forecast drivers)
  - 04_DCF (CFI DCF with dates/periods working, linked to 3-statement forecast)
  - 10K_Sources (clickable SEC links)
  - Charts recreated on DCF (Cash Flow) and 3-statement (Revenue)

Preserves native CFI formulas for taxes, Capex links, UFCF, XNPV, etc.
Forces fullCalcOnLoad so Excel shows values instead of blank formula cells.
"""

from __future__ import annotations

import json
import shutil
import zipfile
from copy import copy
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
FICO_DIR = Path(__file__).resolve().parent
ORIG_3 = ROOT / "DCF resources" / "originals" / "CFI_3-Statement-Model-Complete.xlsx"
ORIG_DCF = ROOT / "DCF resources" / "originals" / "CFI_DCF-Model.xlsx"
CACHE = ROOT / "fico-valuation" / "data" / "fico_companyfacts.json"
OUT = FICO_DIR / "FICO_3S_DCF_Joined.xlsx"
USER_AGENT = "FinancialAnalyst user@example.com"
CIK = "0000814547"
COMPANYFACTS = f"https://data.sec.gov/api/xbrl/companyfacts/CIK{CIK}.json"

BLUE = Font(name="Calibri", color="0000FF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
LINK_FILL = PatternFill("solid", fgColor="E2EFDA")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", color="FFFFFF", bold=True)
TITLE = Font(name="Calibri", size=14, bold=True, color="1F4E79")
NOTE = Font(name="Calibri", size=10, italic=True, color="C00000")
LINK_FONT = Font(name="Calibri", color="0563C1", underline="single")

FILINGS = {
    2025: "https://www.sec.gov/Archives/edgar/data/814547/000081454725000030/fico-20250930.htm",
    2024: "https://www.sec.gov/Archives/edgar/data/814547/000162828024045719/fico-20240930.htm",
    2023: "https://www.sec.gov/Archives/edgar/data/814547/000081454723000022/fico-20230930.htm",
}

HIST_COLS = ["E", "F", "G", "H", "I"]
HIST_YEARS = [2021, 2022, 2023, 2024, 2025]
FORECAST_COLS = ["J", "K", "L", "M", "N"]
# Realistic FICO forecast growth path (explicit period)
GROWTHS = [0.12, 0.10, 0.09, 0.08, 0.07]


def is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def set_input(ws, coord, value, linked=False, allow_formula_overwrite=False) -> None:
    cell = ws[coord]
    if is_formula(cell.value) and not allow_formula_overwrite:
        raise RuntimeError(f"Refusing formula overwrite at {coord}: {cell.value}")
    cell.value = value
    cell.font = Font(name="Calibri", color="000000") if linked else BLUE
    cell.fill = LINK_FILL if linked else INPUT_FILL


def set_formula(ws, coord, formula, linked=True) -> None:
    cell = ws[coord]
    cell.value = formula
    if linked:
        cell.fill = LINK_FILL
        cell.font = Font(name="Calibri", color="000000")


def widen(ws, start=4, end=15, width=16.0) -> None:
    for col in range(start, end + 1):
        letter = get_column_letter(col)
        cur = ws.column_dimensions[letter].width or 10
        ws.column_dimensions[letter].width = max(cur, width)
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 12, 32)


def annual_usd_thousands(facts: dict, concept: str) -> dict[int, float]:
    usgaap = facts["facts"].get("us-gaap", {})
    if concept not in usgaap or "USD" not in usgaap[concept]["units"]:
        return {}
    best: dict[int, dict] = {}
    for it in usgaap[concept]["units"]["USD"]:
        if it.get("form") not in ("10-K", "10-K/A") or it.get("fp") != "FY":
            continue
        end = it.get("end") or ""
        if not end.endswith("-09-30"):
            continue
        y = int(end[:4])
        if y < 2020 or y > 2025:
            continue
        prev = best.get(y)
        if prev is None or it.get("filed", "") >= prev["filed"]:
            best[y] = {"val": it["val"], "filed": it.get("filed", "")}
    return {y: best[y]["val"] / 1000.0 for y in best}


def load_series() -> dict:
    if CACHE.exists() and CACHE.stat().st_size > 1000:
        facts = json.loads(CACHE.read_text())
    else:
        import urllib.request

        req = urllib.request.Request(
            COMPANYFACTS, headers={"User-Agent": USER_AGENT, "Accept": "application/json"}
        )
        with urllib.request.urlopen(req, timeout=60) as resp:
            facts = json.loads(resp.read().decode())
        CACHE.parent.mkdir(parents=True, exist_ok=True)
        CACHE.write_text(json.dumps(facts))

    s = {
        "revenue": annual_usd_thousands(facts, "RevenueFromContractWithCustomerExcludingAssessedTax"),
        "cogs": annual_usd_thousands(facts, "CostOfRevenue"),
        "sga": annual_usd_thousands(facts, "SellingGeneralAndAdministrativeExpense"),
        "rd": annual_usd_thousands(facts, "ResearchAndDevelopmentExpense"),
        "opinc": annual_usd_thousands(facts, "OperatingIncomeLoss"),
        "interest": annual_usd_thousands(facts, "InterestExpense"),
        "tax": annual_usd_thousands(facts, "IncomeTaxExpenseBenefit"),
        "ni": annual_usd_thousands(facts, "NetIncomeLoss"),
        "cash": annual_usd_thousands(facts, "CashAndCashEquivalentsAtCarryingValue"),
        "ar": annual_usd_thousands(facts, "AccountsReceivableNetCurrent"),
        "ppe": annual_usd_thousands(facts, "PropertyPlantAndEquipmentNet"),
        "ap": annual_usd_thousands(facts, "AccountsPayableCurrent"),
        "re": annual_usd_thousands(facts, "RetainedEarningsAccumulatedDeficit"),
        "da": annual_usd_thousands(facts, "DepreciationDepletionAndAmortization"),
        "debt_lt": annual_usd_thousands(facts, "LongTermDebtNoncurrent"),
        "debt_cur": annual_usd_thousands(facts, "LongTermDebtCurrent"),
    }
    s["interest"].update({2023: 95_546, 2024: 105_638, 2025: 133_647})
    debt = {y: s["debt_lt"].get(y, 0) + s["debt_cur"].get(y, 0) for y in range(2020, 2026)}
    debt[2024] = 2_209_021
    debt[2025] = 3_055_691
    s["debt"] = debt
    s["capex"] = dict(annual_usd_thousands(facts, "PaymentsToAcquirePropertyPlantAndEquipment"))
    s["capex"].update({2023: 4_237, 2024: 8_884 + 16_667, 2025: 8_922 + 30_485})
    s["market"] = {
        "price": 1046.23,
        "shares_k": 21_597.635,
        "debt_10q": 5_582_389,
        "cash_10q": 248_444 + 56_093,
        "tax_rate": 150_649 / 802_595,
    }
    return s


def inject_three_statement(ws, s: dict) -> None:
    set_input(ws, "E2", HIST_YEARS[0])

    for col, y in zip(HIST_COLS, HIST_YEARS):
        set_input(ws, f"{col}24", round(s["revenue"][y], 1))
        set_input(ws, f"{col}25", round(s["cogs"][y], 1))
        set_input(ws, f"{col}28", round(s["sga"][y], 1))
        set_input(ws, f"{col}29", round(s["rd"][y], 1))
        set_input(ws, f"{col}30", round(s["da"][y], 1))
        set_input(ws, f"{col}31", round(s["interest"][y], 1))
        set_input(ws, f"{col}35", round(s["tax"][y], 1))
        set_input(ws, f"{col}41", round(s["cash"][y], 1))
        set_input(ws, f"{col}42", round(s["ar"][y], 1))
        set_input(ws, f"{col}43", 0)
        set_input(ws, f"{col}44", round(s["ppe"][y], 1))
        set_input(ws, f"{col}48", round(s["ap"][y], 1))
        set_input(ws, f"{col}49", round(s["debt"][y], 1))

        assets_simple = s["cash"][y] + s["ar"][y] + s["ppe"][y]
        liab_simple = s["ap"][y] + s["debt"][y]
        equity_plug = assets_simple - liab_simple
        re = s["re"][y]
        set_input(ws, f"{col}53", round(re, 1))
        set_input(ws, f"{col}52", round(equity_plug - re, 1))

        if y == HIST_YEARS[0]:
            prev_nwc = s["ar"][y - 1] - s["ap"][y - 1]
            prev_debt = s["debt"].get(y - 1, 739_435)
            prev_cash = s["cash"][y - 1]
            open_ppe = s["ppe"][y - 1]
        else:
            prev_nwc = s["ar"][y - 1] - s["ap"][y - 1]
            prev_debt = s["debt"][y - 1]
            prev_cash = s["cash"][y - 1]
            open_ppe = s["ppe"][y - 1]

        nwc = s["ar"][y] - s["ap"][y]
        dnwc = nwc - prev_nwc
        debt_issue = s["debt"][y] - prev_debt
        capex = abs(s["capex"][y])
        set_input(ws, f"{col}64", round(dnwc, 1))
        set_input(ws, f"{col}68", round(capex, 1))
        set_input(ws, f"{col}72", round(debt_issue, 1))
        cfo_proxy = s["ni"][y] + s["da"][y] - dnwc
        target_delta = s["cash"][y] - prev_cash
        equity_issue = target_delta - (cfo_proxy - capex + debt_issue)
        set_input(ws, f"{col}73", round(equity_issue, 1))

        if col == "E":
            set_input(ws, "D78", round(s["cash"][y - 1], 1))
        else:
            set_input(ws, f"{col}77", round(prev_cash, 1))

        set_input(ws, f"{col}85", round(s["ar"][y], 1))
        set_input(ws, f"{col}86", 0)
        set_input(ws, f"{col}87", round(s["ap"][y], 1))
        da = s["da"][y]
        set_input(ws, f"{col}92", round(open_ppe, 1))
        set_input(ws, f"{col}93", round(s["ppe"][y] - open_ppe + da, 1))
        set_input(ws, f"{col}94", round(da, 1))
        set_input(ws, f"{col}98", round(prev_debt, 1))
        set_input(ws, f"{col}99", round(debt_issue, 1))
        set_input(ws, f"{col}101", round(s["interest"][y], 1))

    # Forecast drivers so 3-statement rolls FICO-scale numbers into DCF links
    y0 = 2025
    sga = s["sga"][y0]
    rd = s["rd"][y0]
    capex0 = s["capex"][y0]
    tax_rate = round(s["market"]["tax_rate"], 4)
    da_rate = max(0.005, s["da"][y0] / s["ppe"].get(y0 - 1, s["ppe"][y0]))

    for i, col in enumerate(FORECAST_COLS):
        g = GROWTHS[i]
        set_input(ws, f"{col}7", g)
        set_input(ws, f"{col}13", tax_rate)
        # Absolute opex assumptions (CFI rows 9/10 feed forecast Salaries/Rent)
        set_input(ws, f"{col}9", round(sga * (1 + g) ** (i + 1), 1))
        set_input(ws, f"{col}10", round(rd * (1 + g) ** (i + 1), 1))
        set_input(ws, f"{col}11", round(min(0.35, da_rate), 4))
        set_input(ws, f"{col}18", round(capex0 * (1 + 0.08) ** (i + 1), 1))
        set_input(ws, f"{col}19", 0)
        set_input(ws, f"{col}20", 0)

    # Keep AR/Inv/AP days reasonable for software
    for col in FORECAST_COLS:
        set_input(ws, f"{col}15", 90)  # AR days
        set_input(ws, f"{col}16", 0)  # Inv days
        set_input(ws, f"{col}17", 30)  # AP days

    ws["B131"] = (
        "FICO joined model: historicals from 10-K; forecast drivers yellow. "
        "DCF sheet pulls EBIT/D&A/Capex/ΔNWC from forecast columns J–N."
    )
    ws["B131"].font = NOTE
    widen(ws)


def inject_dcf(ws, s: dict) -> None:
    m = s["market"]
    # Entry at FY2025 close; first explicit period = FY2026 (E19=1) so dates/yearfrac work
    set_input(ws, "D5", round(m["tax_rate"], 4))
    set_input(ws, "D6", 0.096)
    set_input(ws, "D7", 0.03)
    set_input(ws, "D8", 25)
    set_input(ws, "D9", datetime(2025, 9, 30))  # Transaction / valuation date
    set_input(ws, "D10", datetime(2025, 9, 30))  # FYE anchor
    set_input(ws, "D11", m["price"])
    set_input(ws, "D12", m["shares_k"])
    set_input(ws, "D13", m["debt_10q"])
    set_input(ws, "D14", m["cash_10q"])
    set_input(ws, "D15", round(s["capex"][2025], 1))

    # Write concrete dates/years so time periods always DISPLAY (not blank before calc).
    # Explicit period starts at FY2026 (t=1) so YEARFRAC from entry 2025-09-30 works.
    # Keep YEARFRAC / tax / UFCF / XNPV as native formulas.
    for i, col in enumerate(["E", "F", "G", "H", "I"]):
        year = 2026 + i
        set_input(ws, f"{col}17", year, allow_formula_overwrite=True)
        set_input(ws, f"{col}18", datetime(year, 9, 30), allow_formula_overwrite=True)
        ws[f"{col}18"].number_format = "yyyy-mm-dd"
        set_input(ws, f"{col}19", i + 1, allow_formula_overwrite=True)

    ws["B2"] = "FICO DCF (linked to 03_Three_Statement)"
    ws["B50"] = (
        "Green cells are live links from 03_Three_Statement forecast (J–N). "
        "Yellow = assumptions. Open Excel and allow calculation if values look blank."
    )
    ws["B50"].font = NOTE
    widen(ws, width=16)


def link_dcf_to_three_statement(dcf_ws) -> None:
    """Roll 3-statement forecast into DCF explicit period columns E–I."""
    # Tax rate from first forecast year assumption
    dcf_ws["D5"] = "='03_Three_Statement'!J13"
    dcf_ws["D5"].fill = LINK_FILL
    dcf_ws["D5"].font = Font(name="Calibri", color="000000")

    # Capex assumption fallback still in D15 (used if needed); per-year Capex linked below
    for dcol, scol in zip(["E", "F", "G", "H", "I"], FORECAST_COLS):
        # EBIT ≈ Gross Profit - Salaries - Rent/Overhead - D&A
        set_formula(
            dcf_ws,
            f"{dcol}21",
            (
                f"='03_Three_Statement'!{scol}26-'03_Three_Statement'!{scol}28"
                f"-'03_Three_Statement'!{scol}29-'03_Three_Statement'!{scol}30"
            ),
        )
        # Keep native tax formula (=EBIT * tax rate) — already in template
        set_formula(dcf_ws, f"{dcol}23", f"='03_Three_Statement'!{scol}30")
        set_formula(dcf_ws, f"{dcol}24", f"='03_Three_Statement'!{scol}68")
        set_formula(dcf_ws, f"{dcol}25", f"='03_Three_Statement'!{scol}89")

    dcf_ws["B21"] = "EBIT (from 3-statement forecast)"
    dcf_ws["B22"] = "Less: Cash Taxes (formula)"
    dcf_ws["B23"] = "Plus: D&A (from 3-statement)"
    dcf_ws["B24"] = "Less: Capex (from 3-statement)"
    dcf_ws["B25"] = "Less: Changes in NWC (from 3-statement)"

    # Ensure tax / UFCF / year-fraction formulas still present
    for col in ["E", "F", "G", "H", "I"]:
        if not is_formula(dcf_ws[f"{col}22"].value):
            dcf_ws[f"{col}22"] = f"={col}21*$D$5"
        if not is_formula(dcf_ws[f"{col}26"].value):
            dcf_ws[f"{col}26"] = f"={col}21-{col}22+{col}23-{col}24-{col}25"
        if not is_formula(dcf_ws[f"{col}20"].value):
            prev = "D" if col == "E" else chr(ord(col) - 1)
            dcf_ws[f"{col}20"] = f"=YEARFRAC({prev}18,{col}18)"
        # Transaction CF formulas
        if not is_formula(dcf_ws[f"{col}28"].value):
            dcf_ws[f"{col}28"] = f"=({col}27+{col}26)*{col}20"


def add_charts(dcf_ws, three_ws) -> None:
    # Remove broken leftover charts then add clean ones
    dcf_ws._charts = []
    three_ws._charts = []

    # DCF: Unlevered FCF by year (row 26, cols E-I); cats = year headers row 17
    chart = BarChart()
    chart.type = "col"
    chart.title = "Unlevered Free Cash Flow (linked)"
    chart.y_axis.title = "USD thousands"
    chart.x_axis.title = "Forecast year"
    data = Reference(dcf_ws, min_col=5, min_row=26, max_col=9, max_row=26)
    cats = Reference(dcf_ws, min_col=5, min_row=17, max_col=9, max_row=17)
    chart.add_data(data, from_rows=True, titles_from_data=False)
    chart.set_categories(cats)
    chart.shape = 4
    chart.width = 18
    chart.height = 10
    dcf_ws.add_chart(chart, "L3")

    # 3-statement: Revenue history + forecast
    rev = LineChart()
    rev.title = "FICO Revenue (history E–I, forecast J–N)"
    rev.y_axis.title = "USD thousands"
    rev.x_axis.title = "Fiscal year"
    rev_data = Reference(three_ws, min_col=5, min_row=24, max_col=14, max_row=24)
    rev_cats = Reference(three_ws, min_col=5, min_row=2, max_col=14, max_row=2)
    rev.add_data(rev_data, from_rows=True, titles_from_data=False)
    rev.set_categories(rev_cats)
    rev.width = 18
    rev.height = 10
    three_ws.add_chart(rev, "P3")


def copy_sheet(src_ws, dst_ws) -> None:
    for row in src_ws.iter_rows():
        for cell in row:
            new = dst_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                new.font = copy(cell.font)
                new.fill = copy(cell.fill)
                new.border = copy(cell.border)
                new.number_format = cell.number_format
                new.alignment = copy(cell.alignment)
    for letter, dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[letter].width = dim.width
    for idx, dim in src_ws.row_dimensions.items():
        dst_ws.row_dimensions[idx].height = dim.height
    for merged in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(merged))


def add_sources_sheet(wb, s: dict) -> None:
    ws = wb.create_sheet("10K_Sources", 0)
    ws["A1"] = "FICO joined 3-Statement + DCF — 10-K sources"
    ws["A1"].font = TITLE
    ws["A2"] = "Green DCF cells pull from 03_Three_Statement. Yellow = inputs. Units $ thousands."
    headers = [
        "Item",
        "Where in model",
        "FICO 10-K label",
        "Naming note",
        "FY2025",
        "FY2025 10-K URL",
    ]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(4, c, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    rows = [
        ("Revenue", "03_Three_Statement!I24", "Revenues", "Same", s["revenue"][2025]),
        ("COGS", "03_Three_Statement!I25", "Cost of revenues", "CFI=COGS; 10-K=Cost of revenues", s["cogs"][2025]),
        ("SG&A", "03_Three_Statement!I28", "Selling, general and administrative", "CFI 'Salaries' ← FICO SG&A", s["sga"][2025]),
        ("R&D / OpEx piece", "03_Three_Statement!I29", "Research and development", "CFI 'Rent/Overhead' ← FICO R&D; OpEx≈SG&A+R&D", s["rd"][2025]),
        ("EBIT", "04_DCF!E21 (linked)", "Operating income", "DCF EBIT from 3-stmt GP−SG&A−R&D−D&A", s["opinc"][2025]),
        ("D&A", "03_Three_Statement!I30 → DCF", "Depreciation and amortization", "Cash flow add-back", s["da"][2025]),
        ("CapEx", "03_Three_Statement!J68:N68 → DCF", "PPE purchases + capitalized software", "FY25 8,922+30,485=39,407", s["capex"][2025]),
        ("ΔNWC", "03_Three_Statement!J89:N89 → DCF", "Derived Δ(AR−AP)", "Not a single printed total", ""),
        ("Interest", "03_Three_Statement!I31", "Interest expense, net", "Use net IS line", s["interest"][2025]),
        ("Tax", "03_Three_Statement!I35 / J13", "Provision for income taxes", "Rate linked into DCF!D5", s["tax"][2025]),
        ("Debt (DCF bridge)", "04_DCF!D13", "10-Q total debt bridge", "Interim override of YE 10-K debt", s["market"]["debt_10q"]),
        ("Cash (DCF bridge)", "04_DCF!D14", "10-Q cash + marketable securities", "Interim bridge", s["market"]["cash_10q"]),
    ]
    for i, (item, where, label, note, val) in enumerate(rows, start=5):
        ws.cell(i, 1, item)
        ws.cell(i, 2, where)
        ws.cell(i, 3, label)
        ws.cell(i, 4, note)
        ws.cell(i, 5, val if val != "" else None)
        link = ws.cell(i, 6, FILINGS[2025])
        link.hyperlink = FILINGS[2025]
        link.font = LINK_FONT

    ws["A20"] = "Filing links"
    ws["A21"] = "FY2025 10-K"
    ws["B21"] = FILINGS[2025]
    ws["B21"].hyperlink = FILINGS[2025]
    ws["B21"].font = LINK_FONT
    ws["A22"] = "FY2024 10-K"
    ws["B22"] = FILINGS[2024]
    ws["B22"].hyperlink = FILINGS[2024]
    ws["B22"].font = LINK_FONT
    ws["A23"] = "FY2023 10-K"
    ws["B23"] = FILINGS[2023]
    ws["B23"].hyperlink = FILINGS[2023]
    ws["B23"].font = LINK_FONT
    ws["A25"] = (
        "If Date / Cash Taxes / Capex / UFCF look blank: Excel has not calculated yet. "
        "Press F9 or enable automatic calculation. fullCalcOnLoad is set on this workbook."
    )
    ws["A25"].font = NOTE

    for i, w in enumerate([22, 36, 40, 55, 14, 70], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_instructions(wb) -> None:
    ws = wb.create_sheet("01_Instructions", 0)
    ws["A1"] = "FICO Joined 3-Statement + DCF Model"
    ws["A1"].font = TITLE
    lines = [
        "",
        "This workbook merges the CFI 3-statement and CFI DCF templates into one file.",
        "",
        "Sheet map",
        "  01_Instructions — this page",
        "  10K_Sources — SEC 10-K links + naming notes (OpEx = SG&A + R&D, etc.)",
        "  03_Three_Statement — FICO FY2021–FY2025 history + forecast J–N",
        "  04_DCF — valuation; green cells are FORMULAS linked to 03_Three_Statement",
        "",
        "How data rolls 3-statement → DCF",
        "  04_DCF!E21:I21  EBIT      ← 03 GP − Salaries − Rent/Overhead − D&A (J26:N26 etc.)",
        "  04_DCF!E23:I23  D&A       ← 03_Three_Statement!J30:N30",
        "  04_DCF!E24:I24  Capex     ← 03_Three_Statement!J68:N68",
        "  04_DCF!E25:I25  ΔNWC      ← 03_Three_Statement!J89:N89",
        "  04_DCF!D5       Tax rate  ← 03_Three_Statement!J13",
        "  Cash Taxes / UFCF / XNPV remain native CFI formulas (auto-calc).",
        "",
        "Time periods",
        "  Entry / valuation date D9 = 2025-09-30",
        "  First explicit period E19 = 1 → FY2026 … FY2030 (so dates & year fractions populate)",
        "",
        "Charts",
        "  04_DCF has Unlevered FCF bar chart",
        "  03_Three_Statement has Revenue line chart",
        "",
        "Verify",
        "  03_Three_Statement!I24 = 1,990,869 (FY25 revenue)",
        "  04_DCF!E18 shows a date (FY2026), E19=1, E22/E24/E26 are formulas",
        "  04_DCF!E21 begins with ='03_Three_Statement'!",
    ]
    for i, line in enumerate(lines, start=2):
        ws[f"A{i}"] = line
    ws.column_dimensions["A"].width = 110


def force_full_calc(wb) -> None:
    wb.calculation = CalcProperties(fullCalcOnLoad=True, calcMode="auto", forceFullCalc=True)


def build() -> Path:
    if not ORIG_3.exists() or not ORIG_DCF.exists():
        raise SystemExit("Missing pristine CFI originals in DCF resources/originals/")

    s = load_series()
    print("FY25 revenue", s["revenue"][2025], "capex", s["capex"][2025])

    tmp3 = FICO_DIR / "_tmp_3.xlsx"
    tmpd = FICO_DIR / "_tmp_dcf.xlsx"
    shutil.copy2(ORIG_3, tmp3)
    shutil.copy2(ORIG_DCF, tmpd)

    wb3 = load_workbook(tmp3, data_only=False)
    wbd = load_workbook(tmpd, data_only=False)
    # Drop cover images only (they break re-save); keep sheet structure
    for wb in (wb3, wbd):
        for ws in wb.worksheets:
            if "Cover" in ws.title and getattr(ws, "_images", None):
                ws._images = []

    inject_three_statement(wb3["3 Statement Model"], s)
    inject_dcf(wbd["DCF Model"], s)

    # Assemble joined workbook
    out = Workbook()
    out.remove(out.active)
    add_instructions(out)
    add_sources_sheet(out, s)

    ws3 = out.create_sheet("03_Three_Statement")
    copy_sheet(wb3["3 Statement Model"], ws3)
    wsd = out.create_sheet("04_DCF")
    copy_sheet(wbd["DCF Model"], wsd)

    link_dcf_to_three_statement(wsd)
    add_charts(wsd, ws3)
    force_full_calc(out)

    # Order sheets
    order = ["01_Instructions", "10K_Sources", "03_Three_Statement", "04_DCF"]
    for i, name in enumerate(order):
        out.move_sheet(name, offset=i - out.sheetnames.index(name))

    out.save(OUT)
    tmp3.unlink(missing_ok=True)
    tmpd.unlink(missing_ok=True)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")
    return OUT


def verify(path: Path) -> None:
    wb = load_workbook(path, data_only=False)
    assert set(["01_Instructions", "10K_Sources", "03_Three_Statement", "04_DCF"]).issubset(set(wb.sheetnames))
    s = wb["03_Three_Statement"]
    d = wb["04_DCF"]
    assert s["I24"].value == 1990869 or s["I24"].value == 1990869.0
    assert is_formula(s["J24"].value)
    assert d["E19"].value == 1
    assert d["E17"].value == 2026
    assert d["I17"].value == 2030
    assert isinstance(d["E18"].value, datetime)
    assert is_formula(d["E21"].value) and "03_Three_Statement" in d["E21"].value
    assert is_formula(d["E22"].value)
    assert is_formula(d["E24"].value) and "03_Three_Statement" in d["E24"].value
    assert is_formula(d["E26"].value)
    assert is_formula(d["D32"].value) and "XNPV" in d["D32"].value
    assert len(d._charts) >= 1
    assert len(s._charts) >= 1
    assert wb.calculation.fullCalcOnLoad is True
    print("VERIFY OK")
    print("  3S I24 revenue =", s["I24"].value)
    print("  DCF years E17:I17 =", [d[f"{c}17"].value for c in "EFGHI"])
    print("  DCF dates E18:I18 =", [d[f"{c}18"].value for c in "EFGHI"])
    print("  DCF E21 =", d["E21"].value)
    print("  DCF E24 Capex link =", d["E24"].value)
    print("  DCF charts =", len(d._charts), "3S charts =", len(s._charts))


if __name__ == "__main__":
    p = build()
    verify(p)
