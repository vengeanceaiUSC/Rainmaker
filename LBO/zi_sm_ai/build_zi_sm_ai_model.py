#!/usr/bin/env python3
"""Build ZoomInfo S&M → AI SDR EBITDA margin expansion workbook from FY2025 10-K."""

from __future__ import annotations

import csv
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent / "output" / "ZI_SM_AI_EBITDA_Margin_Model.xlsx"
REPO = ROOT.parent.parent / "ZI_SM_AI_EBITDA_Margin_Model.xlsx"
CSV = ROOT / "zi_fy2025_10k_extract.csv"

BLUE = Font(name="Calibri", color="0000FF")
BLACK = Font(name="Calibri", color="000000")
BOLD = Font(name="Calibri", bold=True)
TITLE = Font(name="Calibri", size=16, bold=True, color="1F4E79")
HDR_F = Font(name="Calibri", color="FFFFFF", bold=True)
HDR = PatternFill("solid", fgColor="1F4E79")
IN = PatternFill("solid", fgColor="FFF2CC")
CALC = PatternFill("solid", fgColor="E2EFDA")
SEC = PatternFill("solid", fgColor="D6DCE4")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
LINK = Font(name="Calibri", color="0563C1", underline="single")


def lab(ws, r, c, t, bold=False):
    cell = ws.cell(r, c, t)
    cell.font = BOLD if bold else BLACK
    return cell


def inp(ws, r, c, v, fmt=None):
    cell = ws.cell(r, c, v)
    cell.font = BLUE
    cell.fill = IN
    cell.border = THIN
    if fmt:
        cell.number_format = fmt
    return cell


def calc(ws, r, c, v, fmt=None):
    cell = ws.cell(r, c, v)
    cell.font = BLACK
    cell.fill = CALC
    cell.border = THIN
    if fmt:
        cell.number_format = fmt
    return cell


def build() -> Path:
    wb = Workbook()

    c = wb.active
    c.title = "Cover"
    c["B2"] = "ZoomInfo (GTM/ZI) — S&M → AI SDR Margin Expansion"
    c["B2"].font = TITLE
    c["B3"] = "Educational ops / LBO model from FY2025 Form 10-K. Open in Excel to recalculate."
    c["B5"] = "Thesis"
    c["C5"] = "Replace outbound SDRs with AI agents; expand EBITDA margin ≥500 bps"
    c["B6"] = "Issuer"
    c["C6"] = "ZoomInfo Technologies Inc. (ticker GTM / legacy ZI)"
    c["B7"] = "CIK"
    c["C7"] = "0001794515"
    c["B8"] = "Source 10-K"
    c["C8"] = "https://www.sec.gov/Archives/edgar/data/1794515/000179451526000012/zi-20251231.htm"
    c["C8"].hyperlink = c["C8"].value
    c["C8"].font = LINK
    c["B9"] = "Filed"
    c["C9"] = "2026-02-12 (FY ended 2025-12-31)"
    c["B10"] = "Companyfacts"
    c["C10"] = "https://data.sec.gov/api/xbrl/companyfacts/CIK0001794515.json"
    c["C10"].hyperlink = c["C10"].value
    c["C10"].font = LINK
    c["B12"] = "Sheets"
    for i, t in enumerate(
        [
            "01_10K_Source — SEC XBRL + human capital extract",
            "02_SM_Isolation — S&M bloat vs target % of revenue",
            "03_Headcount_AI — SDR reduction schedule + AI agent ramp",
            "04_EBITDA_Bridge — margin expansion to 500+ bps",
            "05_Sensitivity — SDR cut % × AI cost → bps grid",
        ],
        start=13,
    ):
        c[f"B{i}"] = t
    c["B19"] = "Disclaimer"
    c["C19"] = "SDR % of S&M and AI unit costs are yellow assumptions. Illustrative — not investment advice."
    c.column_dimensions["B"].width = 22
    c.column_dimensions["C"].width = 100

    s = wb.create_sheet("01_10K_Source")
    s["B2"] = "FY2025 10-K source financials ($ millions unless noted)"
    s["B2"].font = TITLE
    for col, t in enumerate(["Line item", "FY2023", "FY2024", "FY2025", "XBRL / note"], start=2):
        cell = s.cell(4, col, t)
        cell.fill = HDR
        cell.font = HDR_F

    rows = [
        (5, "Revenue", 1239.5, 1214.3, 1249.5, "RevenueFromContractWithCustomerExcludingAssessedTax"),
        (6, "Selling & Marketing", 408.5, 414.1, 414.6, "SellingAndMarketingExpense"),
        (7, "Research & Development", 191.5, 196.1, 182.0, "ResearchAndDevelopmentExpense"),
        (8, "General & Administrative", 179.6, 295.3, 206.7, "GeneralAndAdministrativeExpense"),
        (9, "Operating expenses (ex-COGS)", None, None, 824.2, "OperatingExpenses"),
        (10, "Operating income", 259.5, 97.4, 225.7, "OperatingIncomeLoss"),
        (11, "Depreciation", 19.6, 25.9, 30.3, "Depreciation"),
        (12, "Amort. of intangibles", 61.0, 59.8, 58.5, "AmortizationOfIntangibleAssets"),
        (13, "Total D&A (dep+intang)", 80.6, 85.7, 88.8, "OtherDepreciationAndAmortization"),
        (14, "Share-based compensation", 167.6, 138.0, 116.2, "ShareBasedCompensation"),
        (15, "Advertising expense", None, None, 37.3, "AdvertisingExpense"),
        (16, "Amort. deferred commissions", 75.3, 67.6, 90.2, "AmortizationOfDeferredSalesCommissions"),
        (17, "Net income", 107.3, 29.1, 124.2, "NetIncomeLoss"),
    ]
    for r, name, y23, y24, y25, note in rows:
        lab(s, r, 2, name)
        if y23 is not None:
            inp(s, r, 3, y23, "#,##0.0")
        if y24 is not None:
            inp(s, r, 4, y24, "#,##0.0")
        inp(s, r, 5, y25, "#,##0.0")
        lab(s, r, 6, note)

    lab(s, 19, 2, "Implied COGS (Rev − OpEx − OpInc)", True)
    calc(s, 19, 5, "=E5-E9-E10", "#,##0.0")
    lab(s, 20, 2, "GAAP EBITDA proxy (OpInc + D&A)", True)
    calc(s, 20, 5, "=E10+E13", "#,##0.0")
    lab(s, 21, 2, "GAAP EBITDA margin")
    calc(s, 21, 5, "=E20/E5", "0.0%")
    lab(s, 22, 2, "Adj. EBITDA proxy (+SBC)", True)
    calc(s, 22, 5, "=E20+E14", "#,##0.0")
    lab(s, 23, 2, "Adj. EBITDA margin")
    calc(s, 23, 5, "=E22/E5", "0.0%")

    lab(s, 25, 2, "HUMAN CAPITAL (10-K Item 1)", True)
    s["B25"].fill = SEC
    lab(s, 26, 2, "Total employees (12/31/2025)")
    inp(s, 26, 5, 3180, "#,##0")
    lab(s, 27, 2, "Cost of service employees")
    inp(s, 27, 5, 662, "#,##0")
    lab(s, 28, 2, "Sales & Marketing employees")
    inp(s, 28, 5, 1370, "#,##0")
    lab(s, 29, 2, "R&D employees")
    inp(s, 29, 5, 763, "#,##0")
    lab(s, 30, 2, "G&A employees")
    inp(s, 30, 5, 385, "#,##0")
    lab(s, 31, 2, "Headcount check (sum)")
    calc(s, 31, 5, "=E27+E28+E29+E30", "#,##0")
    lab(s, 32, 2, "S&M $ / S&M employee ($000s)")
    calc(s, 32, 5, "=E6*1000/E28", "#,##0.0")
    lab(s, 33, 2, "S&M % of revenue")
    calc(s, 33, 5, "=E6/E5", "0.0%")

    lab(s, 35, 2, "Quote — Human Capital")
    s["B36"] = (
        "As of December 31, 2025, we had 3,180 employees, consisting of 662 in cost of service, "
        "1,370 in sales and marketing, 763 in research and development, and 385 in general and administrative."
    )
    s.merge_cells("B36:F38")
    s["B36"].alignment = Alignment(wrap_text=True)
    for col, w in [("B", 42), ("C", 12), ("D", 12), ("E", 14), ("F", 70)]:
        s.column_dimensions[col].width = w

    m = wb.create_sheet("02_SM_Isolation")
    m["B2"] = "Isolate bloated Sales & Marketing"
    m["B2"].font = TITLE
    lab(m, 4, 2, "FY2025 Revenue ($m)")
    calc(m, 4, 3, "='01_10K_Source'!E5", "#,##0.0")
    lab(m, 5, 2, "FY2025 S&M ($m)")
    calc(m, 5, 3, "='01_10K_Source'!E6", "#,##0.0")
    lab(m, 6, 2, "S&M % revenue")
    calc(m, 6, 3, "=C5/C4", "0.0%")

    lab(m, 8, 2, "S&M STACK (FY2025)", True)
    m["B8"].fill = SEC
    lab(m, 9, 2, "Advertising")
    calc(m, 9, 3, "='01_10K_Source'!E15", "#,##0.0")
    lab(m, 10, 2, "Deferred commission amortization")
    calc(m, 10, 3, "='01_10K_Source'!E16", "#,##0.0")
    lab(m, 11, 2, "Residual S&M (people + other)")
    calc(m, 11, 3, "=C5-C9-C10", "#,##0.0")
    lab(m, 12, 2, "Residual as % of S&M")
    calc(m, 12, 3, "=C11/C5", "0.0%")
    lab(m, 13, 2, "Residual $ / S&M employee ($000s)")
    calc(m, 13, 3, "=C11*1000/'01_10K_Source'!E28", "#,##0.0")

    lab(m, 15, 2, "BENCHMARK / THESIS INPUTS (edit yellow)", True)
    m["B15"].fill = SEC
    lab(m, 16, 2, "Target mature SaaS S&M % rev")
    inp(m, 16, 3, 0.22, "0.0%")
    lab(m, 17, 2, "S&M bloat vs target ($m)")
    calc(m, 17, 3, "=MAX(0,C5-C4*C16)", "#,##0.0")
    lab(m, 18, 2, "Bloat as % of revenue")
    calc(m, 18, 3, "=C17/C4", "0.0%")
    lab(m, 19, 2, "Bloat in basis points")
    calc(m, 19, 3, "=C18*10000", "#,##0")
    m["B21"] = (
        "10-K: 1,370 S&M employees; S&M is ~33% of revenue. Model treats outbound SDR/BDR capacity "
        "as the AI-replaceable slice (yellow % on next sheet)."
    )
    m.column_dimensions["B"].width = 48
    m.column_dimensions["C"].width = 14

    h = wb.create_sheet("03_Headcount_AI")
    h["B2"] = "Headcount reduction: outbound SDRs → AI agents"
    h["B2"].font = TITLE
    lab(h, 4, 2, "BASELINE", True)
    h["B4"].fill = SEC
    lab(h, 5, 2, "S&M employees (10-K)")
    calc(h, 5, 3, "='01_10K_Source'!E28", "#,##0")
    lab(h, 6, 2, "% of S&M that are outbound SDR/BDR")
    inp(h, 6, 3, 0.40, "0.0%")
    lab(h, 7, 2, "Outbound SDR headcount (Y0)")
    calc(h, 7, 3, "=ROUND(C5*C6,0)", "#,##0")
    lab(h, 8, 2, "Fully loaded cost / SDR / yr ($000s)")
    inp(h, 8, 3, 155, "#,##0.0")
    lab(h, 9, 2, "Addressable SDR cost pool ($m)")
    calc(h, 9, 3, "=C7*C8/1000", "#,##0.0")

    lab(h, 11, 2, "AI AGENT UNIT ECONOMICS", True)
    h["B11"].fill = SEC
    lab(h, 12, 2, "AI agent seats per SDR replaced")
    inp(h, 12, 3, 1.0, "0.0")
    lab(h, 13, 2, "AI platform + compute / seat / yr ($000s)")
    inp(h, 13, 3, 15, "#,##0.0")
    lab(h, 14, 2, "Human oversight FTE per 20 AI seats")
    inp(h, 14, 3, 1.0, "0.0")
    lab(h, 15, 2, "Oversight fully loaded / FTE / yr ($000s)")
    inp(h, 15, 3, 150, "#,##0.0")
    lab(h, 16, 2, "One-time severance / SDR ($000s)")
    inp(h, 16, 3, 25, "#,##0.0")
    lab(h, 17, 2, "One-time AI build / integration ($m) in Y1")
    inp(h, 17, 3, 8.0, "#,##0.0")

    lab(h, 19, 2, "REDUCTION SCHEDULE (% of Y0 SDRs exited by year-end)", True)
    h["B19"].fill = SEC
    lab(h, 20, 2, "Year")
    for i, y in enumerate([2026, 2027, 2028]):
        cell = h.cell(20, 3 + i, y)
        cell.fill = HDR
        cell.font = HDR_F
    lab(h, 21, 2, "Cumulative % SDRs replaced")
    inp(h, 21, 3, 0.45, "0%")
    inp(h, 21, 4, 0.80, "0%")
    inp(h, 21, 5, 0.95, "0%")
    lab(h, 22, 2, "SDRs remaining")
    for i, col in enumerate(["C", "D", "E"]):
        calc(h, 22, 3 + i, f"=ROUND($C$7*(1-{col}21),0)", "#,##0")
    lab(h, 23, 2, "SDRs exited (cumulative)")
    for i, col in enumerate(["C", "D", "E"]):
        calc(h, 23, 3 + i, f"=$C$7-{col}22", "#,##0")
    lab(h, 24, 2, "SDRs exited this year")
    calc(h, 24, 3, "=C23", "#,##0")
    calc(h, 24, 4, "=D23-C23", "#,##0")
    calc(h, 24, 5, "=E23-D23", "#,##0")
    lab(h, 25, 2, "Avg SDRs during year")
    calc(h, 25, 3, "=($C$7+C22)/2", "#,##0.0")
    calc(h, 25, 4, "=(C22+D22)/2", "#,##0.0")
    calc(h, 25, 5, "=(D22+E22)/2", "#,##0.0")

    lab(h, 27, 2, "AI SEATS & OVERSIGHT", True)
    h["B27"].fill = SEC
    lab(h, 28, 2, "AI seats (cum.)")
    for i, col in enumerate(["C", "D", "E"]):
        calc(h, 28, 3 + i, f"={col}23*$C$12", "#,##0.0")
    lab(h, 29, 2, "Oversight FTEs")
    for i, col in enumerate(["C", "D", "E"]):
        calc(h, 29, 3 + i, f"=ROUNDUP({col}28/(20/$C$14),0)", "#,##0")

    lab(h, 31, 2, "P&L IMPACT ($m)", True)
    h["B31"].fill = SEC
    lab(h, 32, 2, "SDR cash cost (avg HC × cost)")
    for i, col in enumerate(["C", "D", "E"]):
        calc(h, 32, 3 + i, f"={col}25*$C$8/1000", "#,##0.0")
    lab(h, 33, 2, "Baseline SDR cost if no program")
    for i in range(3):
        calc(h, 33, 3 + i, "=$C$9", "#,##0.0")
    lab(h, 34, 2, "Gross SDR savings vs baseline")
    for i, col in enumerate(["C", "D", "E"]):
        calc(h, 34, 3 + i, f"={col}33-{col}32", "#,##0.0")
    lab(h, 35, 2, "AI seat opex")
    for i, col in enumerate(["C", "D", "E"]):
        calc(h, 35, 3 + i, f"={col}28*$C$13/1000", "#,##0.0")
    lab(h, 36, 2, "Oversight opex")
    for i, col in enumerate(["C", "D", "E"]):
        calc(h, 36, 3 + i, f"={col}29*$C$15/1000", "#,##0.0")
    lab(h, 37, 2, "Severance (this year's exits)")
    for i, col in enumerate(["C", "D", "E"]):
        calc(h, 37, 3 + i, f"={col}24*$C$16/1000", "#,##0.0")
    lab(h, 38, 2, "AI build (Y1 only)")
    calc(h, 38, 3, "=$C$17", "#,##0.0")
    inp(h, 38, 4, 0, "#,##0.0")
    inp(h, 38, 5, 0, "#,##0.0")
    lab(h, 39, 2, "Net EBITDA benefit (in-year, after one-time)", True)
    for i, col in enumerate(["C", "D", "E"]):
        calc(h, 39, 3 + i, f"={col}34-{col}35-{col}36-{col}37-{col}38", "#,##0.0")
    lab(h, 40, 2, "Year-end run-rate benefit (ex one-time)", True)
    # Steady-state at year-end headcount: full-year cost of remaining vs Y0, minus AI+oversight on exited seats
    for i, col in enumerate(["C", "D", "E"]):
        calc(
            h,
            40,
            3 + i,
            f"=({col}23*$C$8/1000)-({col}28*$C$13/1000)-({col}29*$C$15/1000)",
            "#,##0.0",
        )

    h.column_dimensions["B"].width = 48
    for col in range(3, 6):
        h.column_dimensions[get_column_letter(col)].width = 12

    e = wb.create_sheet("04_EBITDA_Bridge")
    e["B2"] = "EBITDA margin expansion bridge (target ≥500 bps)"
    e["B2"].font = TITLE
    lab(e, 4, 2, "FY2025 baseline Revenue ($m)")
    calc(e, 4, 3, "='01_10K_Source'!E5", "#,##0.0")
    lab(e, 5, 2, "FY2025 GAAP EBITDA proxy ($m)")
    calc(e, 5, 3, "='01_10K_Source'!E20", "#,##0.0")
    lab(e, 6, 2, "Baseline EBITDA margin")
    calc(e, 6, 3, "=C5/C4", "0.0%")
    lab(e, 7, 2, "Target margin expansion (bps)")
    inp(e, 7, 3, 500, "#,##0")
    lab(e, 8, 2, "Required incremental EBITDA ($m)")
    calc(e, 8, 3, "=C4*C7/10000", "#,##0.0")

    lab(e, 10, 2, "Year")
    for i, y in enumerate([2026, 2027, 2028]):
        cell = e.cell(10, 3 + i, y)
        cell.fill = HDR
        cell.font = HDR_F
    lab(e, 11, 2, "Revenue (flat case for margin math)")
    for i in range(3):
        calc(e, 11, 3 + i, "=$C$4", "#,##0.0")
    lab(e, 12, 2, "Net EBITDA benefit from program")
    calc(e, 12, 3, "='03_Headcount_AI'!C39", "#,##0.0")
    calc(e, 12, 4, "='03_Headcount_AI'!D39", "#,##0.0")
    calc(e, 12, 5, "='03_Headcount_AI'!E39", "#,##0.0")
    lab(e, 13, 2, "Run-rate benefit (ex one-time)")
    calc(e, 13, 3, "='03_Headcount_AI'!C40", "#,##0.0")
    calc(e, 13, 4, "='03_Headcount_AI'!D40", "#,##0.0")
    calc(e, 13, 5, "='03_Headcount_AI'!E40", "#,##0.0")
    lab(e, 14, 2, "Pro forma EBITDA")
    for i, col in enumerate(["C", "D", "E"]):
        calc(e, 14, 3 + i, f"=$C$5+{col}12", "#,##0.0")
    lab(e, 15, 2, "Pro forma EBITDA margin")
    for i, col in enumerate(["C", "D", "E"]):
        calc(e, 15, 3 + i, f"={col}14/{col}11", "0.0%")
    lab(e, 16, 2, "Margin expansion (bps)", True)
    for i, col in enumerate(["C", "D", "E"]):
        calc(e, 16, 3 + i, f"=({col}15-$C$6)*10000", "#,##0.0")
    lab(e, 17, 2, "Run-rate expansion (bps)", True)
    for i, col in enumerate(["C", "D", "E"]):
        calc(e, 17, 3 + i, f"=({col}13/$C$4)*10000", "#,##0.0")
    lab(e, 18, 2, "Hits ≥500 bps run-rate?")
    for i, col in enumerate(["C", "D", "E"]):
        calc(e, 18, 3 + i, f'=IF({col}17>=$C$7,"YES","NO")')

    lab(e, 20, 2, "S&M % revenue bridge")
    lab(e, 21, 2, "Baseline S&M %")
    calc(e, 21, 3, "='02_SM_Isolation'!C6", "0.0%")
    lab(e, 22, 2, "PF S&M $ (baseline − SDR save + AI/oversight)")
    for i, col in enumerate(["C", "D", "E"]):
        calc(
            e,
            22,
            3 + i,
            f"='01_10K_Source'!E6-'03_Headcount_AI'!{col}34+'03_Headcount_AI'!{col}35+'03_Headcount_AI'!{col}36",
            "#,##0.0",
        )
    lab(e, 23, 2, "PF S&M % revenue")
    for i, col in enumerate(["C", "D", "E"]):
        calc(e, 23, 3 + i, f"={col}22/$C$4", "0.0%")
    e["B25"] = (
        "Yellow cells on 03_Headcount_AI drive the bridge. Base case sized to clear 500 bps by Y3 run-rate: "
        "40% of S&M = outbound SDRs; 95% replaced; $155k/SDR vs $15k AI seat + oversight. "
        "500 bps on $1,249.5m revenue needs ~$62.5m run-rate EBITDA."
    )
    e.column_dimensions["B"].width = 55
    for col in range(3, 6):
        e.column_dimensions[get_column_letter(col)].width = 12

    x = wb.create_sheet("05_Sensitivity")
    x["B2"] = "Sensitivity — Y3 run-rate margin expansion (bps)"
    x["B2"].font = TITLE
    x["B3"] = "Rows = cumulative % SDRs replaced; Cols = AI $/seat-yr ($000s). Other inputs from 03_Headcount_AI."
    lab(x, 5, 2, "SDR0")
    calc(x, 5, 3, "='03_Headcount_AI'!C7", "#,##0")
    lab(x, 6, 2, "SDR cost $k")
    calc(x, 6, 3, "='03_Headcount_AI'!C8", "#,##0.0")
    lab(x, 7, 2, "Seats/SDR")
    calc(x, 7, 3, "='03_Headcount_AI'!C12", "0.0")
    lab(x, 8, 2, "Oversight $k")
    calc(x, 8, 3, "='03_Headcount_AI'!C15", "#,##0.0")
    lab(x, 9, 2, "Revenue $m")
    calc(x, 9, 3, "='01_10K_Source'!E5", "#,##0.0")

    pcts = [0.5, 0.6, 0.7, 0.8, 0.9, 1.0]
    ai_costs = [8, 12, 18, 24, 30, 40]
    x["B11"] = "bps \\ AI $k"
    for j, a in enumerate(ai_costs):
        inp(x, 11, 3 + j, a, "#,##0")
    for i, p in enumerate(pcts):
        inp(x, 12 + i, 2, p, "0%")
        for j, a in enumerate(ai_costs):
            calc(
                x,
                12 + i,
                3 + j,
                f"=(($C$5*{p})*$C$6/1000-($C$5*{p})*$C$7*{a}/1000-ROUNDUP(($C$5*{p})*$C$7/20,0)*$C$8/1000)/$C$9*10000",
                "#,##0",
            )
    lab(x, 20, 2, "Base case Y3 run-rate bps (from bridge)")
    calc(x, 20, 3, "='04_EBITDA_Bridge'!E17", "#,##0.0")
    x.column_dimensions["B"].width = 18
    for col in range(3, 9):
        x.column_dimensions[get_column_letter(col)].width = 10

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    wb.save(REPO)

    with CSV.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["metric", "fy2023", "fy2024", "fy2025", "source"])
        w.writerow(["revenue", 1239.5, 1214.3, 1249.5, "XBRL RevenueFromContractWithCustomerExcludingAssessedTax"])
        w.writerow(["selling_and_marketing", 408.5, 414.1, 414.6, "XBRL SellingAndMarketingExpense"])
        w.writerow(["operating_income", 259.5, 97.4, 225.7, "XBRL OperatingIncomeLoss"])
        w.writerow(["da_total", 80.6, 85.7, 88.8, "XBRL OtherDepreciationAndAmortization"])
        w.writerow(["sbc", 167.6, 138.0, 116.2, "XBRL ShareBasedCompensation"])
        w.writerow(["advertising", "", "", 37.3, "XBRL AdvertisingExpense"])
        w.writerow(["commission_amort", 75.3, 67.6, 90.2, "XBRL AmortizationOfDeferredSalesCommissions"])
        w.writerow(["employees_total", "", "", 3180, "10-K Human Capital"])
        w.writerow(["employees_sm", "", "", 1370, "10-K Human Capital"])

    # Sanity: Y3 end-state run-rate bps (matches default yellow inputs)
    sdr0 = round(1370 * 0.40)
    exited = round(sdr0 * 0.95)
    save = exited * 155 / 1000
    ai = exited * 15 / 1000
    ov = math.ceil(exited / 20) * 150 / 1000
    bps = (save - ai - ov) / 1249.5 * 10000
    print(f"Wrote {OUT}")
    print(f"Wrote {REPO}")
    print(f"Base Y3 end-state approx: SDR0={sdr0}, exited={exited}, net=${save-ai-ov:.1f}m, bps≈{bps:.0f}")
    return OUT


if __name__ == "__main__":
    build()
