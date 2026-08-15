#!/usr/bin/env python3
"""
Rebuild Assumptions_Drivers so EVERY row has:
  B Driver | C Value | D Why reasonable (user commentary) |
  E Source link | F Source credibility note | G Feeds (clickable jump) | H Secondary source
Remove any AI_Cost_Sources tab. Make Feeds hyperlinks jump to the math cells.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

REPO = Path(__file__).resolve().parents[2]
XLSX = REPO / "vengeanceaiUSC_LBOMODEL2.xlsx"
XLSX2 = REPO / "LBO" / "output" / "vengeanceaiUSC_LBOMODEL2.xlsx"
PDF = REPO / "LBO" / "output" / "LBOMODEL2_DRIVER_SOURCES.pdf"
MD = REPO / "LBO" / "output" / "LBOMODEL2_DRIVER_SOURCES.md"
NOTES = REPO / "LBO" / "output" / "LBOMODEL2_NOTES.md"

RAW_XLSX = "https://raw.githubusercontent.com/vengeanceaiUSC/Rainmaker/cursor/vengeanceaiusclbomodel2-44dc/vengeanceaiUSC_LBOMODEL2.xlsx"
RAW_PDF = "https://raw.githubusercontent.com/vengeanceaiUSC/Rainmaker/cursor/vengeanceaiusclbomodel2-44dc/LBO/output/LBOMODEL2_DRIVER_SOURCES.pdf"

HDR = PatternFill("solid", fgColor="1F4E79")
HDR_F = Font(name="Calibri", color="FFFFFF", bold=True, size=10)
TITLE = Font(name="Calibri", size=14, bold=True, color="1F4E79")
BOLD = Font(name="Calibri", bold=True)
BLACK = Font(name="Calibri", size=10)
BLUE = Font(name="Calibri", color="0000FF", size=10)
LINK = Font(name="Calibri", color="0563C1", underline="single", size=9)
NOTE = PatternFill("solid", fgColor="E7F3FF")
IN = PatternFill("solid", fgColor="FFF2CC")
CALC = PatternFill("solid", fgColor="E2EFDA")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
WRAP = Alignment(wrap_text=True, vertical="top")

# row: name, value, fmt, user_why, source_url, source_label, credibility_30w, feed_label, feed_jump, secondary_url
# feed_jump = "Sheet!Cell" for hyperlink target
DRIVERS = [
    (
        5,
        "Revenue growth",
        0.10,
        "0.0%",
        "Assumes steady top-line expansion. Reasonable for a mature enterprise software company, avoiding aggressive hypergrowth assumptions while maintaining a healthy, realistic sales trajectory.",
        "https://www.businesswire.com/news/home/20250225694335/en/ZoomInfo-Announces-Fourth-Quarter-and-Full-Year-2024-Financial-Results",
        "ZoomInfo FY2024 results (BusinessWire)",
        "ZI disclosed FY24 revenue of $1,214.3m (−2% YoY). Using 10% forward growth is a mature-SaaS recovery underwrite held equal in AI vs Base—not hypergrowth.",
        "Go to AI_Operating!D5 (Revenue)",
        "AI_Operating!D5",
        "https://valueaddvc.com/blog/saas-evebitda-multiples-explained-benchmarks-and-what-they-mean",
    ),
    (
        6,
        "Gross margin",
        0.80,
        "0.0%",
        "High profitability on direct services. Highly reasonable for a SaaS company, which historically scales digital infrastructure with minimal incremental cost per user.",
        "https://www.getaleph.com/answers/saas-gross-margin-2026",
        "Aleph/Benchmarkit SaaS GM median ~80%",
        "Independent SaaS benchmarks put software median GM near 80%; ZI already prints ~84% GAAP GM, so 80% is conservative and industry-standard.",
        "Go to AI_Operating!D6 (COGS)",
        "AI_Operating!D6",
        "https://www.readyratios.com/sec/GTM_zoominfo-technologies-inc",
    ),
    (
        7,
        "Sales payroll cut (Y1)",
        0.15,
        "0.0%",
        "Immediate headcount reduction for human SDRs. Reasonable as a private equity cost-cutting measure, easily achieved by eliminating the lower-tier outbound sales team.",
        "https://firstsales.io/blog/cost-per-meeting-outbound/",
        "AI SDR vs human SDR cost/meeting analysis",
        "Published outbound economics show fully loaded human SDRs far costlier than AI stacks. −15% Y1 payroll is a modest PE-style cut of outbound headcount, not a full S&M wipeout.",
        "Go to AI_Operating!D7 (Sales payroll)",
        "AI_Operating!D7",
        "https://www.sec.gov/Archives/edgar/data/1794515/000179451525000045/zi-20241231.htm",
    ),
    (
        8,
        "Variable commission cut",
        0.20,
        "0.0%",
        "Reduces payouts for automated sales. Reasonable because AI agents do not collect commission checks, retaining more profit per enterprise deal directly for the company.",
        "https://syncgtm.com/blog/how-to-pay-sales-development-rep",
        "SDR pay mix ~70/30 base/variable (SyncGTM)",
        "SaaS SDR OTE is typically ~70% base / ~30% variable. Replacing humans with AI shrinks the commission pool; −20% is a partial cut of that variable layer.",
        "Go to AI_Operating!D8 (Commissions)",
        "AI_Operating!D8",
        "https://syncgtm.com/blog/how-much-do-sales-development-representatives-make-at-tech-companies",
    ),
    (
        9,
        "AI infrastructure ($m)",
        15.0,
        "#,##0.0",
        "Fixed cost for AI software and compute. Reasonable estimate for enterprise-scale AI deployment, covering API calls and infrastructure needed to replace human reps.",
        "https://www.salesforce.com/news/press-releases/2025/05/15/agentforce-flexible-pricing-news/",
        "Salesforce Agentforce $2/conv + Flex Credits press",
        "SF publishes $2/conversation and ~$0.10/action. At ZI scale (~378 outbound roles = 25% of 1,513 S&M), $5–10k/mo agents imply ~$23–45m; $15m is conservative.",
        "Go to AI_Operating!D10 (AI infrastructure)",
        "AI_Operating!D10",
        "https://highticketaisystems.com/blog/ai-sdr-pricing-explained",
    ),
    (
        10,
        "S&M expense growth Y2+",
        0.02,
        "0.0%",
        "Caps future marketing cost increases. Reasonable because AI scales infinitely without needing proportional headcount additions, severing the link between revenue growth and human hiring.",
        "https://www.gartner.com/en/newsroom/press-releases/2026-08-10-gartner-forecasts-worldwide-artificial-intelligence-optimized-iaas-spending-to-grow-96-percent-in-2026",
        "Gartner: agentic AI / inference scales as opex",
        "Gartner documents production agentic AI as inference opex that scales without linear headcount. Cap S&M growth at 2% to model that break from revenue-linked hiring.",
        "Go to AI_Operating!E7 (payroll Y2+)",
        "AI_Operating!E7",
        "https://www.ciodive.com/news/AI-spending-soars-enterprise-maturity/827488/",
    ),
    (
        11,
        "G&A % of revenue",
        0.18,
        "0.0%",
        "General administrative overhead costs. Reasonable for a public-to-private SaaS company, maintaining standard corporate expenses without relying on unrealistic operational magic outside of sales.",
        "https://www.businesswire.com/news/home/20250225694335/en/ZoomInfo-Announces-Fourth-Quarter-and-Full-Year-2024-Financial-Results",
        "ZoomInfo FY24 G&A $295.3m (company results)",
        "ZI FY24 G&A was $295.3m (~24% of revenue). Locking ~18% keeps standard corporate overhead while the thesis only attacks the sales floor—not finance/comp magic.",
        "Go to AI_Operating!D12 (G&A)",
        "AI_Operating!D12",
        "https://www.sec.gov/Archives/edgar/data/1794515/000179451525000045/zi-20241231.htm",
    ),
    (
        12,
        "CapEx reduction vs base",
        0.05,
        "0.0%",
        "Lowers physical equipment spending. Reasonable because firing 15% of the sales team permanently eliminates the need to purchase and refresh their laptops and hardware.",
        "https://caliberoutbound.com/blog/build-vs-buy-outbound-the-2026-cost-math",
        "In-house SDR all-in includes devices/tooling",
        "Human SDR all-in budgets include laptops and seats. Cutting outbound humans removes refresh CapEx; −5% vs base CapEx rate is a small, directionally correct haircut.",
        "Go to AI_Operating!D19 (CapEx)",
        "AI_Operating!D19",
        "https://firstsales.io/blog/cost-per-meeting-outbound/",
    ),
    (
        13,
        "WC / receivables improvement",
        0.10,
        "0.0%",
        "Speeds up cash collection from clients. Reasonable because automated AI billing systems can aggressively follow up on late invoices, freeing up cash for debt paydown.",
        "https://www.readyratios.com/sec/GTM_zoominfo-technologies-inc",
        "ZI receivables ~80 days (SEC-derived ratios)",
        "ZI receivables ~80 days show WC is real. Automated billing follow-ups can shorten DSO; −10% WC intensity vs base is a modest collections gain, not a wipeout.",
        "Go to AI_Operating!D20 (ΔNWC)",
        "AI_Operating!D20",
        "https://www.sec.gov/Archives/edgar/data/1794515/000179451525000045/zi-20241231.htm",
    ),
    (
        14,
        "SOFR",
        0.043,
        "0.00%",
        "The foundational secured lending interest rate. Reasonable and historically accurate for the macroeconomic environment, properly reflecting actual institutional baseline borrowing costs.",
        "https://www.newyorkfed.org/markets/reference-rates/sofr",
        "NY Fed official SOFR reference rate",
        "SOFR is the NY Fed’s official overnight Treasury-repo financing benchmark used to price leveraged loans. ~4.3% is a planning level near recent SOFR prints (see FRED).",
        "Go to Debt_Sweep_AI!C12 (TLA rate uses SOFR)",
        "Debt_Sweep_AI!C12",
        "https://fred.stlouisfed.org/series/SOFR",
    ),
    (
        15,
        "Term Loan A spread",
        0.04,
        "0.00%",
        "The risk premium for senior debt. Reasonable for a standard LBO, reflecting typical syndicated commercial bank pricing for secured corporate borrowing.",
        "https://ryanoconnellfinance.com/lbo-debt-structure/",
        "LBO debt structure: TLA illustrative spreads",
        "Bank TLA in LBO guides typically prices tighter than TLB (often SOFR+~275–375bps). SOFR+400bps is a conservative senior bank spread independent of AI ops.",
        "Go to Assumptions_Drivers!C27 (TLA rate)",
        "Assumptions_Drivers!C27",
        "https://ibinterviewquestions.com/guides/debt-capital-markets/term-loan-b-tlb-mechanics-and-why-it-dominates-lev-lending",
    ),
    (
        16,
        "Term Loan B spread",
        0.05,
        "0.00%",
        "The risk premium for riskier debt. Reasonable, as Term Loan B is usually held by institutional investors requiring higher yields for slightly higher risk.",
        "https://ibinterviewquestions.com/guides/valuation-investment-banking/lbo-debt-structures-senior-subordinated-mezzanine",
        "TLB pricing often SOFR+300–500bps",
        "Institutional TLB commonly prices SOFR+300–500bps. SOFR+500bps is the upper published band—set by syndicated loan markets, not by internal AI strategy.",
        "Go to Assumptions_Drivers!C28 (TLB rate)",
        "Assumptions_Drivers!C28",
        "https://ryanoconnellfinance.com/lbo-debt-structure/",
    ),
    (
        17,
        "Mandatory amortization",
        0.01,
        "0.0%",
        "Required annual principal repayment. Highly reasonable and standard market practice for leveraged loans, forcing a tiny minimum debt reduction each year.",
        "https://ibinterviewquestions.com/guides/debt-capital-markets/term-loan-b-tlb-mechanics-and-why-it-dominates-lev-lending",
        "Syndicated TLB ~1% annual amort standard",
        "Market-standard institutional TLB schedules ~1% annual amortization with a bullet maturity. Loan docs—not AI FCF—set this mandatory floor.",
        "Go to Debt_Sweep_AI!D5 (Mandatory 1%)",
        "Debt_Sweep_AI!D5",
        "https://clearvaluelending.com/glossary/term-loan-b",
    ),
    (
        18,
        "Cash sweep %",
        1.0,
        "0.0%",
        "Dedicates all excess cash to debt. Reasonable and standard for maximizing LBO returns, ensuring every free dollar generated immediately aggressively reduces outstanding leverage.",
        "https://ryanoconnellfinance.com/lbo-debt-structure/",
        "LBO optional prepay / excess cash practice",
        "Sponsor LBOs typically sweep excess FCF to optional debt prepay after mandatory amort. 100% sweep maximizes deleveraging so AI margins hit MOIC/IRR.",
        "Go to Debt_Sweep_AI!E5 (Total paydown)",
        "Debt_Sweep_AI!E5",
        "https://ibinterviewquestions.com/guides/valuation-investment-banking/lbo-debt-structures-senior-subordinated-mezzanine",
    ),
    (
        19,
        "Tax rate",
        0.21,
        "0.0%",
        "Corporate income tax obligation. Perfectly reasonable because it strictly matches the current standard U.S. federal statutory corporate tax rate without aggressive tax-dodging assumptions.",
        "https://taxsummaries.pwc.com/united-states/corporate/taxes-on-corporate-income",
        "PwC Tax Summaries: US federal CIT 21%",
        "US federal corporate income tax is a flat 21% statutory rate. Operating-margin gains do not change the federal statutory rate used here.",
        "Go to AI_Operating!D18 (Net income)",
        "AI_Operating!D18",
        "https://tradingeconomics.com/united-states/corporate-tax-rate",
    ),
    (
        20,
        "Exit EV/EBITDA multiple",
        13.0,
        "0.0x",
        "The valuation multiple upon selling. Reasonable to keep it conservative; relying on a 13x exit assumes no market inflation, proving returns come from operational growth.",
        "https://valueaddvc.com/blog/saas-evebitda-multiples-explained-benchmarks-and-what-they-mean",
        "PE SaaS buyouts often ~15–22x EBITDA",
        "PE SaaS deals often underwrite ~15–22x EBITDA. Locking exit at 13x in BOTH cases is deliberately conservative so returns come from ops, not multiple expansion.",
        "Go to LBO!D17 (Exit multiple)",
        "LBO!D17",
        "https://aventis-advisors.com/software-valuation-multiples/",
    ),
    (
        21,
        "S&M baseline ($m)",
        425.0,
        "#,##0.0",
        "This is historical assumptions without AI. Total initial sales spend is reasonable, representing roughly 35% of FY24 revenue, perfectly aligning with standard SaaS benchmarks.",
        "https://www.businesswire.com/news/home/20250225694335/en/ZoomInfo-Announces-Fourth-Quarter-and-Full-Year-2024-Financial-Results",
        "ZI FY24 S&M $414.1m on $1,214.3m revenue",
        "Company-reported FY24 S&M = $414.1m (~34% of $1,214.3m). Model ~$425m (~35%) is a rounded baseline consistent with that filing.",
        "Go to AI_Operating!C9 (S&M human FY24A)",
        "AI_Operating!C9",
        "https://www.sec.gov/Archives/edgar/data/1794515/000179451525000045/zi-20241231.htm",
    ),
    (
        22,
        "Payroll baseline ($m)",
        318.8,
        "#,##0.0",
        "This is historical assumptions without AI. The fixed salary portion is reasonable, representing 75% of S&M, accurately reflecting that human software sales floors are base-salary heavy.",
        "https://syncgtm.com/blog/how-to-pay-sales-development-rep",
        "SDR OTE mostly base (~70%) — SyncGTM",
        "With ~70/30 base/variable SDR mix, most people cost is payroll. Payroll baseline = 75% of S&M matches that base-salary-heavy sales-floor structure.",
        "Go to AI_Operating!C7 (Payroll FY24A)",
        "AI_Operating!C7",
        "https://pulserevops.com/knowledge/ra0197",
    ),
    (
        23,
        "Commission baseline ($m)",
        106.2,
        "#,##0.0",
        "This is historical assumptions without AI. The variable payout is reasonable, calculating exactly to the remaining 25% of the budget, representing standard quota-based enterprise sales incentives.",
        "https://syncgtm.com/blog/how-much-do-sales-development-representatives-make-at-tech-companies",
        "Variable ~30–35% of tech SDR OTE",
        "Tech SDR variable is typically ~30–35% of OTE. Setting commission pool at 25% of S&M is a slightly conservative carve of total S&M for variable payouts.",
        "Go to AI_Operating!C8 (Commissions FY24A)",
        "AI_Operating!C8",
        "https://syncgtm.com/blog/how-to-pay-sales-development-rep",
    ),
    (
        24,
        "CapEx base % of rev",
        0.02,
        "0.0%",
        "This is historical assumptions without AI. Normal hardware and capitalized software spending is reasonable for asset-light tech companies investing heavily in code rather than factories.",
        "https://www.businesswire.com/news/home/20250225694335/en/ZoomInfo-Announces-Fourth-Quarter-and-Full-Year-2024-Financial-Results",
        "ZI FY24 strong FCF / asset-light SaaS profile",
        "SaaS is asset-light vs industrials; ZI still printed strong FY24 FCF. 2% of revenue is a simple maintenance CapEx planning rate before the AI −5% haircut.",
        "Go to AI_Operating!D19 (CapEx)",
        "AI_Operating!D19",
        "https://www.sec.gov/Archives/edgar/data/1794515/000179451525000045/zi-20241231.htm",
    ),
    (
        25,
        "WC base % of Δrev",
        0.02,
        "0.0%",
        "This is historical assumptions without AI. Normal working capital needs are reasonable, showing that as revenue grows, a small percentage of cash gets tied up in daily operations.",
        "https://www.readyratios.com/sec/GTM_zoominfo-technologies-inc",
        "ZI receivables days ~80 (WC intensity)",
        "Receivables ~80 days imply WC scales with growth. Using 2% of Δrevenue as ΔNWC is a simplified SaaS WC plug anchored to that receivables reality.",
        "Go to AI_Operating!D20 (ΔNWC)",
        "AI_Operating!D20",
        "https://www.sec.gov/Archives/edgar/data/1794515/000179451525000045/zi-20241231.htm",
    ),
    (
        27,
        "TLA rate (SOFR+4)",
        None,  # formula
        "0.00%",
        "Total Term Loan A interest. Mathematically correct and reasonable for the total borrowing cost demanded by senior lenders in this specific rate environment.",
        "https://ryanoconnellfinance.com/lbo-debt-structure/",
        "LBO TLA = SOFR + senior bank spread",
        "Computed as SOFR (C14) + TLA spread (C15). All-in ~8.3% matches senior secured bank pricing in this SOFR + 400bps setup.",
        "Go to Debt_Sweep_AI!C12 (TLA rate)",
        "Debt_Sweep_AI!C12",
        "https://www.newyorkfed.org/markets/reference-rates/sofr",
    ),
    (
        28,
        "TLB rate (SOFR+5)",
        None,
        "0.00%",
        "Total Term Loan B interest. Mathematically correct and perfectly reasonable, pricing the institutional debt risk appropriately above the safer Term Loan A.",
        "https://ibinterviewquestions.com/guides/valuation-investment-banking/lbo-debt-structures-senior-subordinated-mezzanine",
        "LBO TLB = SOFR + institutional spread",
        "Computed as SOFR (C14) + TLB spread (C16). All-in ~9.3% prices institutional TLB risk appropriately above TLA in this structure.",
        "Go to Debt_Sweep_AI!D12 (TLB rate)",
        "Debt_Sweep_AI!D12",
        "https://www.newyorkfed.org/markets/reference-rates/sofr",
    ),
]


def add_ext_link(cell, url, label):
    cell.value = label
    cell.hyperlink = url
    cell.font = LINK
    cell.alignment = WRAP
    cell.border = THIN


def add_jump(cell, sheet_cell, label):
    """Internal workbook jump: Sheet!A1"""
    sht, addr = sheet_cell.split("!")
    addr0 = addr.split(":")[0]
    cell.value = label
    cell.hyperlink = f"#'{sht}'!{addr0}"
    cell.font = LINK
    cell.alignment = WRAP
    cell.border = THIN


def rebuild_drivers(wb):
    if "AI_Cost_Sources" in wb.sheetnames:
        del wb["AI_Cost_Sources"]

    # Recreate Assumptions_Drivers cleanly
    # Preserve values by rewriting in place if exists, else create
    if "Assumptions_Drivers" in wb.sheetnames:
        # delete and recreate at front-ish position
        idx = wb.sheetnames.index("Assumptions_Drivers")
        del wb["Assumptions_Drivers"]
    else:
        idx = 1
    d = wb.create_sheet("Assumptions_Drivers", min(idx, 1))

    d["B2"] = "Assumptions Drivers — EVERY row: Value + Why reasonable + Source link + clickable Feeds"
    d["B2"].font = TITLE
    d["B3"] = (
        "Columns on the SAME row as each driver: D = why reasonable (model commentary); "
        "E = clickable source URL; F = ~30-word source credibility; G = clickable Feeds jump into the model. "
        "No separate AI cost tab — all evidence lives here."
    )
    d["B3"].alignment = WRAP
    d.merge_cells("B3:H3")
    d.row_dimensions[3].height = 40

    headers = [
        (2, "Driver"),
        (3, "Value"),
        (4, "Why this number is reasonable (model commentary)"),
        (5, "Source link (click)"),
        (6, "Source credibility / why we can use it (~30 words)"),
        (7, "Feeds (click to jump)"),
        (8, "Secondary source (click)"),
    ]
    for col, h in headers:
        cell = d.cell(4, col, h)
        cell.fill = HDR
        cell.font = HDR_F
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    d.row_dimensions[4].height = 36

    for (
        row,
        name,
        val,
        fmt,
        user_why,
        url,
        url_label,
        cred,
        feed_label,
        feed_jump,
        url2,
    ) in DRIVERS:
        d.cell(row, 2, name).font = BLACK
        d.cell(row, 2).border = THIN

        vcell = d.cell(row, 3)
        if row in (27, 28):
            vcell.value = "=C14+C15" if row == 27 else "=C14+C16"
            vcell.fill = CALC
        else:
            vcell.value = val
            vcell.fill = IN
            vcell.font = BLUE
        vcell.number_format = fmt
        vcell.border = THIN

        why = d.cell(row, 4, user_why)
        why.alignment = WRAP
        why.border = THIN
        why.fill = NOTE
        why.font = BLACK

        add_ext_link(d.cell(row, 5), url, url_label)

        cred_c = d.cell(row, 6, cred)
        cred_c.alignment = WRAP
        cred_c.border = THIN
        cred_c.fill = NOTE
        cred_c.font = BLACK

        add_jump(d.cell(row, 7), feed_jump, feed_label)
        add_ext_link(d.cell(row, 8), url2, url2.replace("https://", "").replace("http://", "")[:70])

        d.row_dimensions[row].height = 58

    # spacer label for computed rates
    d["B26"] = "Computed rates (from SOFR + spreads above)"
    d["B26"].font = BOLD

    d["B30"] = "PDF of this source map:"
    d["B30"].font = BOLD
    add_ext_link(d["C30"], RAW_PDF, RAW_PDF)
    d["B31"] = "Educational / research use only — not investment advice."
    d["B31"].font = Font(name="Calibri", italic=True, size=9)

    widths = {
        "B": 28,
        "C": 10,
        "D": 55,
        "E": 40,
        "F": 48,
        "G": 36,
        "H": 40,
    }
    for col, w in widths.items():
        d.column_dimensions[col].width = w
    d.freeze_panes = "B5"

    # Move sheet to front after Strategy if present
    # Keep near front for visibility
    try:
        wb.move_sheet(d, offset=-len(wb.sheetnames))
    except Exception:
        pass
    # Put Assumptions_Drivers as sheet index 0 for visibility
    sheets = wb._sheets
    if d in sheets:
        sheets.remove(d)
        sheets.insert(0, d)

    return d


def scrub_refs(wb):
    for name in list(wb.sheetnames):
        ws = wb[name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and "AI_Cost_Sources" in cell.value:
                    cell.value = cell.value.replace("AI_Cost_Sources", "Assumptions_Drivers cols D–H")
                tgt = getattr(cell.hyperlink, "target", None) if cell.hyperlink else None
                if tgt and "AI_Cost_Sources" in str(tgt):
                    cell.hyperlink = "#'Assumptions_Drivers'!E9"
                    cell.value = "Go to Assumptions_Drivers!E9"
                    cell.font = LINK

    if "Strategy_Summary" in wb.sheetnames:
        s = wb["Strategy_Summary"]
        s["B22"] = "ALL driver sources + commentary live on Assumptions_Drivers (cols D–H, same row). No separate AI cost tab."
        add_jump(s["B23"], "Assumptions_Drivers!D5", "Go to Assumptions_Drivers (start at Revenue growth)")
        add_ext_link(s["B24"], RAW_PDF, RAW_PDF)


def write_pdf():
    styles = getSampleStyleSheet()
    title = ParagraphStyle("T", parent=styles["Heading1"], fontSize=11, spaceAfter=6)
    small = ParagraphStyle("S", parent=styles["Normal"], fontSize=6.5, leading=8)
    body = ParagraphStyle("B", parent=styles["Normal"], fontSize=8, leading=10)
    doc = SimpleDocTemplate(
        str(PDF),
        pagesize=letter,
        leftMargin=0.45 * inch,
        rightMargin=0.45 * inch,
        topMargin=0.4 * inch,
        bottomMargin=0.4 * inch,
    )
    story = [
        Paragraph("LBOMODEL2 — Assumptions_Drivers Source Map (unified; no AI_Cost_Sources tab)", title),
        Paragraph(
            f"Every driver row has Why-reasonable commentary, clickable source, credibility note, and clickable Feeds. "
            f"<link href='{RAW_XLSX}'>Excel</link>",
            body,
        ),
        Spacer(1, 4),
    ]
    data = [["Driver", "Why reasonable", "Source", "Feeds"]]
    for item in DRIVERS:
        _r, name, _v, _f, why, url, label, _cred, feed_label, feed_jump, _u2 = item
        data.append(
            [
                Paragraph(f"<b>{name}</b>", small),
                Paragraph(why, small),
                Paragraph(f'<link href="{url}">{label}</link>', small),
                Paragraph(f"{feed_label}<br/>({feed_jump})", small),
            ]
        )
    t = Table(data, colWidths=[1.2 * inch, 2.6 * inch, 2.0 * inch, 1.6 * inch])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 7),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B0B0B0")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8FC")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    story.append(t)
    story.append(Spacer(1, 6))
    story.append(Paragraph("Educational / research use only — not investment advice.", small))
    doc.build(story)

    lines = [
        "# LBOMODEL2 — Assumptions_Drivers Source Map",
        "",
        "All commentary + links live on **Assumptions_Drivers** (same row). **No** `AI_Cost_Sources` tab.",
        "",
        f"**Excel:** {RAW_XLSX}",
        "",
        f"**PDF:** {RAW_PDF}",
        "",
    ]
    for item in DRIVERS:
        _r, name, _v, _f, why, url, label, cred, feed_label, feed_jump, url2 = item
        lines += [
            f"## {name}",
            f"- **Why reasonable:** {why}",
            f"- **Source:** [{label}]({url})",
            f"- **Credibility:** {cred}",
            f"- **Feeds:** `{feed_jump}` ({feed_label})",
            f"- **Secondary:** {url2}",
            "",
        ]
    MD.write_text("\n".join(lines))
    NOTES.write_text(
        f"""# LBOMODEL2

Open sheet **Assumptions_Drivers** (first tab).

On each driver row:
- **D** — why the number is reasonable (model commentary)
- **E** — clickable primary source
- **F** — ~30-word source credibility
- **G** — clickable Feeds jump into AI_Operating / Debt_Sweep_AI / LBO
- **H** — secondary source

No separate AI cost tab.

Excel: {RAW_XLSX}
PDF: {RAW_PDF}
"""
    )


def main():
    wb = load_workbook(XLSX)
    rebuild_drivers(wb)
    scrub_refs(wb)
    assert "AI_Cost_Sources" not in wb.sheetnames
    wb.save(XLSX)
    wb.save(XLSX2)
    write_pdf()

    wb2 = load_workbook(XLSX)
    d = wb2["Assumptions_Drivers"]
    print("first sheet:", wb2.sheetnames[0])
    print("AI_Cost_Sources:", "AI_Cost_Sources" in wb2.sheetnames)
    n_src = n_feed = n_why = 0
    for r in list(range(5, 26)) + [27, 28]:
        if d.cell(r, 4).value:
            n_why += 1
        if d.cell(r, 5).hyperlink:
            n_src += 1
        if d.cell(r, 7).hyperlink:
            n_feed += 1
        print(
            f"R{r}: why={bool(d.cell(r,4).value)} src={bool(d.cell(r,5).hyperlink)} "
            f"feed={d.cell(r,7).value} -> {d.cell(r,7).hyperlink.target if d.cell(r,7).hyperlink else None}"
        )
    print(f"counts why={n_why} src={n_src} feed={n_feed}")


if __name__ == "__main__":
    main()
