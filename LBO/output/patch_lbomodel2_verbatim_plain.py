#!/usr/bin/env python3
"""
Remove Verbatim HYPERLINK formulas. Columns I/J/K = plain Ctrl+F strings to copy.
Strip #:~:text= from Source E/H links so they just open the page.
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

REPO = Path(__file__).resolve().parents[2]
XLSX = REPO / "vengeanceaiUSC_LBOMODEL2.xlsx"
XLSX2 = REPO / "LBO" / "output" / "vengeanceaiUSC_LBOMODEL2.xlsx"

BLACK = Font(name="Calibri", size=9)
LINK = Font(name="Calibri", color="0563C1", underline="single", size=9)
HDR = PatternFill("solid", fgColor="1F4E79")
HDR_F = Font(name="Calibri", color="FFFFFF", bold=True, size=9)
GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GRAY = PatternFill("solid", fgColor="F2F2F2")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)

# Plain copy-paste Ctrl+F strings (no hyperlinks)
# row -> up to 3 verbatim strings for I/J/K
VERBATIM = {
    5: [
        'Ctrl+F: GAAP Revenue of $1,214.3 million',
        'Ctrl+F: We generated revenue of $1,214.3 million',
        '10% growth: N/A — MODEL CONST',
    ],
    6: [
        'Ctrl+F: The 2025 median software gross margin is 80%',
    ],
    7: [
        'Ctrl+F: down 18% YoY',
    ],
    8: [
        'Ctrl+F: 70/30 base/variable split',
        'Ctrl+F: 30–35% of total pay',
        '20% cut: N/A — MODEL CONST',
    ],
    9: [
        'Ctrl+F: $5,000-10,000+',
        'Ctrl+F: $10,000-$15,000+',
        'Ctrl+F: 1,513 in sales and marketing',
    ],
    10: [
        'Ctrl+F: inflation of 2 percent over the longer run',
    ],
    11: [
        'Ctrl+F: G&A was 18%',
        'Ctrl+F: MEDIAN',
    ],
    12: ['5% CapEx cut: N/A — MODEL CONST'],
    13: ['10% WC improve: N/A — MODEL CONST'],
    14: [
        'Ctrl+F: 4.30',
    ],
    15: [
        'Ctrl+F: SOFR + 400-500 bps',
    ],
    16: [
        'Ctrl+F: SOFR + 300-500 basis points',
    ],
    17: [
        'Ctrl+F: minimal amortization (1% annually)',
    ],
    18: [
        'Ctrl+F: 50-100% of annual excess free cash flow',
        'Ctrl+F: typically modeled at 100%',
    ],
    19: [
        'Ctrl+F: flat rate of 21%',
    ],
    20: [
        'Ctrl+F: 12.9x',
    ],
    21: [
        'Ctrl+F: $1,214.3',
        'Ctrl+F: 414.1',
        'Ctrl+F: $414.1',
    ],
    22: [
        'Ctrl+F: 70/30 base/variable split',
        '$318.8 / 75%: MODEL carve',
    ],
    23: [
        'Ctrl+F: 30–35% of total pay',
        '$106.2 / 25%: MODEL carve',
    ],
    24: [
        'Ctrl+F: CapEx of 1',
        'Tip: page uses en-dash 1–3%; or Ctrl+F: Most pure-play SaaS companies report CapEx of',
        '2%: MODEL midpoint of 1–3% band',
    ],
    25: ['2% WC plug: N/A — MODEL CONST'],
    27: ['See Ctrl+F strings on rows 14–15'],
    28: ['See Ctrl+F strings on rows 14 & 16'],
}

# Clean source URLs (no text fragment) for E/H labels
SOURCE_E = {
    5: (
        "https://markets.financialcontent.com/stocks/article/bizwire-2025-2-25-zoominfo-announces-fourth-quarter-and-full-year-2024-financial-results",
        "IN EQ $ base only: FY24 Rev=$1,214.3m (ZI). C5=10% = MODEL CONST (10% not in filing).",
    ),
    6: (
        "https://www.getaleph.com/answers/saas-gross-margin-2026",
        "IN EQ: software median GM=80% → C6=0.80 (Aleph × Benchmarkit)",
    ),
    7: (
        "https://www.digitalapplied.com/blog/ai-sdr-statistics-2026-outbound-sales-data-points",
        "IN EQ: C7=18% = US B2B SaaS net SDR headcount down 18% YoY (Digital Applied / Bridge Group)",
    ),
    8: (
        "https://syncgtm.com/blog/how-to-pay-sales-development-rep",
        "Pool mix sourced (~30% variable); C8=20% cut = MODEL CONST (20% cut not in source).",
    ),
    9: (
        "https://www.whitespacesolutions.ai/content/ai-sdr-pricing-guide-2026",
        "IN EQ: $5,000 and $10,000/mo AI SDR (White Space + Miniloop) + 1,513 S&M HC (SEC)",
    ),
    10: (
        "https://www.federalreserve.gov/faqs/economy_14400.htm",
        "IN EQ: C10=2% = Fed longer-run inflation goal of 2 percent (FOMC FAQ)",
    ),
    11: (
        "https://blossomstreetventures.medium.com/what-should-saas-spend-on-cogs-s-m-r-d-and-g-a-the-data-from-57-public-cos-8e88eabe99ca",
        "IN EQ: C11=18% = Blossom Street 2025 public SaaS median G&A 18% of revenue",
    ),
    14: (
        "https://www.global-rates.com/en/interest-rates/sofr/historical/2024/",
        "IN EQ: C14=4.30% = SOFR 2024 calendar-year Lowest (Global-Rates summary of NY Fed)",
    ),
    15: (
        "https://ctacquisitions.com/acquisition-financing/",
        "IN EQ: C15=+400bps = floor of senior Term Loan A SOFR+400-500 bps (CT Acquisitions)",
    ),
    16: (
        "https://ibinterviewquestions.com/guides/valuation-investment-banking/lbo-debt-structures-senior-subordinated-mezzanine",
        "IN EQ: TLB SOFR+300-500 bps → C16=+500bps = upper bound",
    ),
    17: (
        "https://clearvaluelending.com/glossary/term-loan-b",
        "IN EQ: ~1%/yr TLB amort → C17=1%",
    ),
    18: (
        "https://ryanoconnellfinance.com/lbo-model-fundamentals/",
        "IN EQ: C18=100% = upper bound of ECF sweep 50-100%; models often use 100%",
    ),
    19: (
        "https://taxsummaries.pwc.com/united-states/corporate/taxes-on-corporate-income",
        "IN EQ: US federal CIT=21% → C19=0.21 (PwC)",
    ),
    20: (
        "https://aventis-advisors.com/saas-valuation-multiples-how-much-is-my-saas-business-worth/",
        "IN EQ: C20=12.9x = Aventis private SaaS EV/EBITDA 1st quartile 12.9x",
    ),
    21: (
        "https://markets.financialcontent.com/stocks/article/bizwire-2025-2-25-zoominfo-announces-fourth-quarter-and-full-year-2024-financial-results",
        "IN EQ: Rev $1,214.3m; S&M $414.1m (34.1%) → model 35%→$425m",
    ),
    22: (
        "https://syncgtm.com/blog/how-to-pay-sales-development-rep",
        "IN EQ: ~70% base OTE → model 75%×$425m=$318.8m",
    ),
    23: (
        "https://syncgtm.com/blog/how-much-do-sales-development-representatives-make-at-tech-companies",
        "IN EQ: residual 25%×$425m=$106.2m (near sourced ~30–35% variable)",
    ),
    24: (
        "https://saasdb.app/learn/financials/fcf-margin/",
        "IN EQ band: SaaS CapEx 1–3% of rev (SaaSDB) → C24=2% = midpoint of that band",
    ),
}

SOURCE_H = {
    5: (
        "https://www.sec.gov/Archives/edgar/data/1794515/000179451525000045/zi-20241231.htm",
        "SEC 10-K same $1,214.3m",
    ),
    8: (
        "https://syncgtm.com/blog/how-much-do-sales-development-representatives-make-at-tech-companies",
        "Also: commission 30–35% of pay",
    ),
    9: (
        "https://www.miniloop.ai/blog/11x-pricing",
        "Miniloop Entry $5,000 / Enterprise $10,000+",
    ),
    11: (
        "https://www.blossomstreetventures.com/saas-metric/2025-annual-percent-of-revenue",
        "Table MEDIAN G&A as % of Revenue = 18%",
    ),
    18: (
        "https://ryanoconnellfinance.com/lbo-model-fundamentals/",
        "Also: typically modeled at 100%",
    ),
    21: (
        "https://www.sec.gov/Archives/edgar/data/1794515/000179451525000045/zi-20241231.htm",
        "SEC 10-K Sales and marketing $414.1",
    ),
}


def esc(s: str) -> str:
    return s.replace('"', '""')


def set_source_link(cell, url: str, label: str):
    """Plain page link via HYPERLINK formula (no text-fragment)."""
    cell.hyperlink = None
    cell.value = f'=HYPERLINK("{esc(url)}","{esc(label)}")'
    cell.font = LINK
    cell.alignment = WRAP


def set_plain(cell, text: str | None, fill):
    cell.hyperlink = None
    cell.value = text
    cell.font = BLACK
    cell.alignment = WRAP
    cell.border = THIN
    cell.fill = fill if text else GRAY


def main():
    wb = load_workbook(XLSX)
    ws = wb["Assumptions_Drivers"]

    ws["I4"] = "Verbatim 1 — copy this into Ctrl+F on the source page"
    ws["J4"] = "Verbatim 2 — copy this into Ctrl+F on the source page"
    ws["K4"] = "Verbatim 3 — copy this into Ctrl+F on the source page"
    for col in ("I", "J", "K"):
        ws[f"{col}4"].fill = HDR
        ws[f"{col}4"].font = HDR_F
        ws[f"{col}4"].alignment = WRAP
        ws.column_dimensions[col].width = 48

    ws["B2"] = (
        "Assumptions Drivers — Col A Baseline Yes/No. "
        "Col E/H = source page link (opens page only). "
        "Cols I/J/K = plain Verbatim strings — copy into Ctrl+F on that source (no hyperlink)."
    )
    ws["B3"] = (
        "How to verify: click Source (E) → open page → copy Verbatim from I/J/K → Ctrl+F → paste."
    )

    for r, parts in VERBATIM.items():
        fill = YELLOW if any("MODEL CONST" in (p or "") or "MODEL carve" in (p or "") for p in parts) else GREEN
        if r in (21, 22, 23):
            fill = GREEN  # baseline sourced pools still have findable strings
        for i, col in enumerate(("I", "J", "K")):
            if i < len(parts):
                set_plain(ws[f"{col}{r}"], parts[i], fill if "N/A" not in parts[i] and "MODEL CONST" not in parts[i] else YELLOW)
            else:
                set_plain(ws[f"{col}{r}"], None, GRAY)
                ws[f"{col}{r}"].value = None

    # Restore E/H as clean page links (no #:~:text=)
    for r, (url, label) in SOURCE_E.items():
        set_source_link(ws[f"E{r}"], url, label)
    for r, (url, label) in SOURCE_H.items():
        set_source_link(ws[f"H{r}"], url, label)

    # Rows with no source link — keep plain E text, clear bad HYPERLINK if any
    for r in (12, 13, 25, 27, 28):
        cell = ws[f"E{r}"]
        if isinstance(cell.value, str) and cell.value.startswith("=HYPERLINK"):
            # extract label roughly or set known plain
            pass
        cell.hyperlink = None
        if r == 12:
            cell.value = "C12=5% CapEx cut = MODEL CONST (no source prints 5%)."
        elif r == 13:
            cell.value = "C13=10% WC improvement = MODEL CONST (no source equals 10%)."
        elif r == 25:
            cell.value = "C25=2% ΔNWC/ΔRev = MODEL CONST (no source equals 2%)."
        elif r == 27:
            cell.value = "EQ only: C27=C14+C15 (see rows 14–15 Verbatim)"
        elif r == 28:
            cell.value = "EQ only: C28=C14+C16 (see rows 14 & 16 Verbatim)"
        cell.font = BLACK
        cell.alignment = WRAP

    ws["B33"] = "HOW TO USE VERBATIM (no auto-highlight):"
    ws["C33"] = (
        "1) Click Source link in column E (opens the page). "
        "2) Copy the string from column I (or J/K). "
        "3) Ctrl+F on the page and paste. "
        "Example row 9: I9 copy $5,000-10,000+ | J9 copy $10,000-$15,000+ | K9 copy 1,513 in sales and marketing"
    )
    ws["C33"].alignment = WRAP
    ws["C33"].font = BLACK

    wb.save(XLSX)
    shutil.copy2(XLSX, XLSX2)

    # Verify
    wb2 = load_workbook(XLSX)
    ws2 = wb2["Assumptions_Drivers"]
    for col in "IJK":
        v = ws2[f"{col}9"].value
        assert v is None or (isinstance(v, str) and not v.startswith("=HYPERLINK")), v
        print(f"{col}9:", v)
    e9 = ws2["E9"].value
    assert "HYPERLINK" in str(e9) and "#:~:text=" not in str(e9), e9
    print("E9 clean:", e9[:80])
    print("OK — Verbatim plain; Source links without text-fragment")


if __name__ == "__main__":
    main()
