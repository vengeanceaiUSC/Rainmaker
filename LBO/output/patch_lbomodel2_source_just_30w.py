#!/usr/bin/env python3
"""Rewrite Assumptions_Drivers col F: Source says X; does X go into the math? ≤30 words, ~10th grade."""
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
PATH = ROOT / "vengeanceaiUSC_LBOMODEL2.xlsx"

F = {
    5: "ZoomInfo says FY24 sales were $1,214.3m. That sales base is in the math. The +10% growth rate is our choice, not from a source.",
    6: "Aleph says the 2025 software median gross margin is 80%. We use that 80% in the COGS math.",
    7: "Digital Applied says U.S. B2B SaaS SDR headcount fell 18% YoY. We use that 18% as the Year 1 payroll cut.",
    8: "SyncGTM says SDR pay is about 70/30 base/variable. That mix is context only. The 20% cut is our choice and goes into the math.",
    9: "White Space/Miniloop say ~$5k–$10k+/mo. ZoomInfo’s 10-K says 1,513 S&M staff. Those go into the $34m math. The 25% role share is our choice.",
    10: "The Fed says longer-run inflation is 2%. We use that 2% to grow S&M costs after Year 1.",
    11: "Blossom Street says 2025 public SaaS median G&A is 18% of sales. We use that 18% in the G&A math.",
    12: "No source says a 5% CapEx cut. The 5% is our choice and still goes into the CapEx math.",
    13: "No source says a 10% working-capital gain. The 10% is our choice and still goes into the WC math.",
    14: "Global-Rates shows the lowest 2024 SOFR print was 4.30%. We use that 4.30% in the interest math.",
    15: "CT Acquisitions lists senior Term Loan A at SOFR + 400–500 bps. We use +400 bps from that range in the rate math.",
    16: "LBO debt guides list Term Loan B at SOFR + 300–500 bps. We use +500 bps from that range in the rate math.",
    17: "ClearValue says Term Loan B amortizes about 1% per year. We use that 1% in the debt paydown math.",
    18: "Ryan O’Connell says excess-cash sweeps are often 50–100% and models often use 100%. We use that 100% in the sweep math.",
    19: "PwC says the U.S. federal corporate tax rate is a flat 21%. We use that 21% in the tax math.",
    20: "Aventis shows private SaaS EV/EBITDA first quartile at 12.9x. We use that 12.9x as the exit multiple.",
    21: "ZoomInfo says FY24 S&M was $414.1m. That exact figure is not the driver. We use $425m (~35% of sales) in the math.",
    22: "SyncGTM says pay is about 70% base. That 70% is not the driver. We use $318.8m (75% of S&M) as a model carve in the math.",
    23: "SyncGTM says commission is about 30–35% of pay. That band is not the driver. We use $106.2m (25% of S&M) as a model carve.",
    24: "SaaSDB says pure-play SaaS CapEx is often 1–3% of sales. We use 2%, the midpoint of that band, in the CapEx math.",
    25: "No source says a 2% WC plug. The 2% is our choice and still goes into the working-capital math.",
    27: "No new source. This cell only adds SOFR (row 14) and the TLA spread (row 15). Both already go into the rate math.",
    28: "No new source. This cell only adds SOFR (row 14) and the TLB spread (row 16). Both already go into the rate math.",
}


def main():
    wb = load_workbook(PATH)
    ws = wb["Assumptions_Drivers"]
    thin = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")
    font = Font(name="Calibri", size=9)
    for r, text in F.items():
        assert len(text.split()) <= 30, (r, len(text.split()), text)
        c = ws.cell(r, 6, text)
        c.font = font
        c.alignment = wrap
        c.border = thin
        c.fill = PatternFill()
    ws.cell(4, 6).value = "Source justification (≤30 words): Source says X. Does X go into the math?"
    ws.cell(4, 6).font = Font(name="Calibri", size=9, bold=True)
    ws.cell(4, 6).alignment = wrap
    ws.column_dimensions["F"].width = 44
    wb.save(PATH)
    print("Updated F justifications")


if __name__ == "__main__":
    main()
